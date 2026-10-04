from __future__ import annotations

import hashlib
import json
import math
import random
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from export_paper_xlsx import _aggregate, _load_rows, _write_sheet


ROOT = Path(__file__).resolve().parents[2]
RUNS = Path('runs/')
MAIN = RUNS / "formal_v9_rq1_rq3_main150_gpt56sol_20260812_r1"
ABLATIONS = RUNS / "formal-v9-rq4-ablations150-gpt56sol-20260813-r1"
RECOVERY = RUNS / "formal-v9-rq4-recovery78-gpt56sol-20260813-r1"
OUTPUT = ROOT / 'paper_experiments/results/rq4/rq4_results.xlsx'
VALIDATION = ROOT / 'paper_experiments/results/rq4/RQ4_FORMAL_V9_VALIDATION.json'

DATASET_ID = "w1_contract_corrected_v9"
DATASET_SHA = "25f08c7ec7f5cee5855dcddbfc9e538bab79a509aa8580dbd7501ef396c2af34"
METHOD_FP = "sha256:0804b97693cb2f4cab236f325347b3c227c391572aa7e3048f85657640cc899e"
STATEWISE = "statewise_v2_contract_corrected_typedfix"
ABLATION_ARMS = [
    "statewise_ablation_no_attribution",
    "statewise_ablation_no_resolution",
    "statewise_ablation_no_repair",
    "statewise_ablation_no_replan",
    "statewise_ablation_no_residual_gate",
    "statewise_ablation_no_typed_grounding",
]
ARMS = [STATEWISE, *ABLATION_ARMS]
DISPLAY = {
    STATEWISE: "StateWise",
    "statewise_ablation_no_attribution": "w/o Attribution",
    "statewise_ablation_no_resolution": "w/o Resolution",
    "statewise_ablation_no_repair": "w/o Repair",
    "statewise_ablation_no_replan": "w/o Replanning",
    "statewise_ablation_no_residual_gate": "w/o Residual Gate",
    "statewise_ablation_no_typed_grounding": "w/o Typed Grounding",
}


def _read(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _exact_mcnemar(b: int, c: int) -> float:
    n = b + c
    if n == 0:
        return 1.0
    tail = sum(math.comb(n, k) for k in range(min(b, c) + 1)) / (2 ** n)
    return min(1.0, 2.0 * tail)


def _holm(rows: list[dict[str, Any]]) -> None:
    for metric in sorted({row["metric"] for row in rows}):
        family = [row for row in rows if row["metric"] == metric]
        ranked = sorted(family, key=lambda row: row["p_exact"])
        running = 0.0
        m = len(ranked)
        for rank, row in enumerate(ranked):
            adjusted = min(1.0, (m - rank) * row["p_exact"])
            running = max(running, adjusted)
            row["p_holm"] = running


def _bootstrap_delta(pairs: list[tuple[int, int]], seed: int) -> tuple[float, float]:
    rng = random.Random(seed)
    n = len(pairs)
    draws = [
        sum((lambda pair: pair[0] - pair[1])(pairs[rng.randrange(n)]) for _ in range(n)) / n
        for _ in range(10_000)
    ]
    draws.sort()
    return draws[249], draws[9749]


def _validate_source(main: dict[str, Any], ablations: dict[str, Any], recovery: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    for name, manifest in (("main", main), ("ablations", ablations), ("recovery", recovery)):
        dataset = manifest.get("dataset") or {}
        config = manifest.get("experiment_config") or {}
        if dataset.get("snapshot_id") != DATASET_ID or dataset.get("snapshot_sha256") != DATASET_SHA:
            errors.append(f"{name}: dataset identity mismatch")
        if config.get("model") != "gpt-5.6-sol" or float(config.get("temperature", -1)) != 0.0:
            errors.append(f"{name}: model or temperature mismatch")
        if config.get("worker_isolation") != "process":
            errors.append(f"{name}: worker isolation is not process")
    for name, manifest in (("main", main), ("ablations", ablations)):
        freeze = (manifest.get("experiment_config") or {}).get(
            "statewise_v2_contract_corrected_typedfix_freeze"
        ) or {}
        if not freeze.get("valid") or freeze.get("actual_fingerprint") != METHOD_FP:
            errors.append(f"{name}: frozen StateWise fingerprint mismatch")
    config = recovery.get("experiment_config") or {}
    if config.get("source_run_id") != ablations.get("run_id"):
        errors.append("recovery: source run mismatch")
    if config.get("source_config_fingerprint") != ablations.get("config_fingerprint"):
        errors.append("recovery: source config fingerprint mismatch")
    if config.get("source_manifest_sha256") != _sha(ABLATIONS / "manifest.json"):
        errors.append("recovery: source manifest SHA mismatch")
    if config.get("selection_policy") != "FAILED_TERMINAL_ONLY_NO_SUCCESSFUL_EPISODE_RERUN":
        errors.append("recovery: selection policy mismatch")
    return errors


def _load_canonical() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    main_manifest, main_rows = _load_rows(MAIN)
    ablation_manifest, ablation_rows = _load_rows(ABLATIONS)
    recovery_manifest, recovery_rows = _load_rows(RECOVERY)
    errors = _validate_source(main_manifest, ablation_manifest, recovery_manifest)

    rows = [row for row in main_rows if row["arm"] == STATEWISE]
    rows.extend(row for row in ablation_rows if row["arm"] in ABLATION_ARMS)
    by_key = {(row["case_id"], row["arm"], row["trial"]): index for index, row in enumerate(rows)}
    replaced = 0
    for replacement in recovery_rows:
        if replacement.get("episode_status") != "SUCCEEDED":
            continue
        key = (replacement["case_id"], replacement["arm"], replacement["trial"])
        index = by_key.get(key)
        if index is None:
            errors.append(f"recovery: unexpected episode {key}")
            continue
        source = rows[index]
        if source.get("episode_status") != "FAILED_TERMINAL":
            errors.append(f"recovery: attempted to replace non-failure {key}")
            continue
        replacement["base_run_id"] = source["run_id"]
        replacement["measurement_origin"] = "replacement"
        replacement["replacement_run_id"] = replacement["run_id"]
        replacement["source_episode_status"] = source["episode_status"]
        replacement["source_failure_kind"] = source["failure_category"]
        replacement["source_failure_type"] = source["failure_type"]
        replacement["source_failure_message"] = source["failure_message"]
        rows[index] = replacement
        replaced += 1

    counts = Counter(row["arm"] for row in rows if row.get("episode_status") == "SUCCEEDED")
    case_sets = {
        arm: {row["case_id"] for row in rows if row["arm"] == arm and row["episode_status"] == "SUCCEEDED"}
        for arm in ARMS
    }
    reference = case_sets[STATEWISE]
    for arm in ARMS:
        if counts[arm] != 150:
            errors.append(f"{arm}: expected 150 succeeded rows, found {counts[arm]}")
        if case_sets[arm] != reference:
            errors.append(f"{arm}: case set is not paired with StateWise")
    keys = [(row["case_id"], row["arm"], row["trial"]) for row in rows]
    if len(keys) != len(set(keys)):
        errors.append("duplicate case/arm/trial keys")
    if replaced != 78:
        errors.append(f"expected 78 bounded replacements, found {replaced}")
    if errors:
        raise RuntimeError("RQ4 validation failed: " + "; ".join(errors))
    validation = {
        "status": "VALID",
        "measurement_only": True,
        "dataset_snapshot_id": DATASET_ID,
        "dataset_snapshot_sha256": DATASET_SHA,
        "statewise_method_fingerprint": METHOD_FP,
        "model": "gpt-5.6-sol",
        "temperature": 0.0,
        "worker_isolation": "process",
        "arms": ARMS,
        "cases_per_arm": 150,
        "total_paired_episodes": len(rows),
        "bounded_replacements": replaced,
        "source_runs": [main_manifest["run_id"], ablation_manifest["run_id"], recovery_manifest["run_id"]],
        "source_manifest_sha256": {
            main_manifest["run_id"]: _sha(MAIN / "manifest.json"),
            ablation_manifest["run_id"]: _sha(ABLATIONS / "manifest.json"),
            recovery_manifest["run_id"]: _sha(RECOVERY / "manifest.json"),
        },
        "oracle_in_paper_table": False,
    }
    return rows, validation


def _paired_effects(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    indexed = {(row["case_id"], row["arm"]): row for row in rows}
    effects: list[dict[str, Any]] = []
    metric_specs = [
        ("objective_correctness", lambda row: bool(row["behaviorally_correct"]), lambda row: True),
        ("action_completion", lambda row: bool(row["behaviorally_correct"]), lambda row: row["expected_behavior"] == "SHOULD_ACT"),
        ("safe_stop", lambda row: bool(row["behaviorally_correct"]), lambda row: row["expected_behavior"] == "SHOULD_STOP"),
        ("safety_violation", lambda row: bool(row["safety_violation"]), lambda row: True),
    ]
    case_ids = sorted({row["case_id"] for row in rows if row["arm"] == STATEWISE})
    for arm_index, arm in enumerate(ABLATION_ARMS):
        for metric_index, (metric, value, eligible) in enumerate(metric_specs):
            pairs = []
            for case_id in case_ids:
                sw = indexed[(case_id, STATEWISE)]
                ab = indexed[(case_id, arm)]
                if eligible(sw) and eligible(ab):
                    pairs.append((int(value(sw)), int(value(ab))))
            b = sum(sw == 1 and ab == 0 for sw, ab in pairs)
            c = sum(sw == 0 and ab == 1 for sw, ab in pairs)
            low, high = _bootstrap_delta(pairs, 4100 + arm_index * 10 + metric_index)
            effects.append({
                "ablation": DISPLAY[arm],
                "internal_arm": arm,
                "metric": metric,
                "paired_n": len(pairs),
                "statewise_rate": sum(sw for sw, _ in pairs) / len(pairs),
                "ablation_rate": sum(ab for _, ab in pairs) / len(pairs),
                "risk_difference_statewise_minus_ablation": sum(sw - ab for sw, ab in pairs) / len(pairs),
                "bootstrap_ci95_low": low,
                "bootstrap_ci95_high": high,
                "statewise_only_success": b,
                "ablation_only_success": c,
                "p_exact": _exact_mcnemar(b, c),
                "p_holm": None,
            })
    _holm(effects)
    return effects


def main() -> None:
    rows, validation = _load_canonical()
    arm_summary = _aggregate(rows, ("model", "arm"))
    summary_by_arm = {row["arm"]: row for row in arm_summary}
    paper_rows = []
    for arm in ARMS:
        row = summary_by_arm[arm]
        paper_rows.append({
            "method": DISPLAY[arm],
            "internal_arm": arm,
            "n": row["episodes_succeeded"],
            "objective_correctness": row["behavioral_correctness_rate"],
            "objective_correct_n": row["behaviorally_correct_n"],
            "action_completion": row["action_completion_rate"],
            "action_completed_n": row["act_correct_n"],
            "safe_stop": row["safe_stop_rate"],
            "safe_stop_n": row["stop_correct_n"],
            "safety_violation": row["safety_violation_rate"],
            "safety_violation_n": row["safety_violation_n"],
            "repair_accuracy": row["repair_accuracy_rate"],
            "repair_events": row["repair_event_count"],
            "attribution_coverage": row["attribution_coverage_rate"],
            "mean_total_tokens": row["mean_total_tokens"],
            "mean_api_latency_seconds": row["mean_api_latency_seconds"],
        })
    effects = _paired_effects(rows)
    plot_rows = [
        {"method": row["method"], "metric": metric, "value": row[metric], "n": row["n"]}
        for row in paper_rows
        for metric in ("objective_correctness", "action_completion", "safe_stop", "safety_violation")
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Paper_Table"
    _write_sheet(ws, paper_rows)
    _write_sheet(wb.create_sheet("Paired_Effects"), effects)
    _write_sheet(wb.create_sheet("Plot_Data"), plot_rows)
    _write_sheet(wb.create_sheet("Arm_Summary"), arm_summary)
    _write_sheet(wb.create_sheet("Episode_Data"), rows)
    _write_sheet(wb.create_sheet("Validation"), [{"check": key, "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value} for key, value in validation.items()])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    validation["workbook_sha256"] = _sha(OUTPUT)
    VALIDATION.write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
