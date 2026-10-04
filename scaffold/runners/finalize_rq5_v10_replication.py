from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean
from typing import Any

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]
RUN_ROOT = Path('runs/')
RESULT_ROOT = ROOT / 'paper_experiments/results/rq5'
PRIMARY = RESULT_ROOT / "rq5_results.xlsx"
OUTPUT = RESULT_ROOT / "rq5_v10_replication3.xlsx"
VALIDATION = RESULT_ROOT / "RQ5_V10_REPLICATION3_VALIDATION.json"

DATASET_SHA = "a871a91427dd4a5bead0115be695bb74df28a660b25990103cdece681b4b7a62"
EXPECTED_CODE_HASHES = {
    "scaffold/runners/run_rq5_longitudinal_real.py": "43d82c8b558d521ce04386c1b1991a173ecc511d941cd087387e1044d369ff7a",
    "scaffold/statewise/evaluation/longitudinal_runtime.py": "22155e13e81af07f558e0901d051f76aed745fe92c4520a4d667e39747e1d811",
    "scaffold/statewise/evaluation/longitudinal_schema.py": "efc90579f490e901ccef46a2c0fb1228f19a67379c6c7350f7dacae2fdd72317",
    "scaffold/statewise/experiments/pilot_conditions.py": "ba12ee8679e7f507e3a117b7f8ff6351bd43e16f9945969542e8e1a79955750a",
    "scaffold/statewise/runtime/lightweight_runtime.py": "844f36ee193f05cb1c12b46a34de55855a58f6b053b9a1171d601d1c314ae4ec",
    "scaffold/statewise/runtime/structured_mutation_runtime.py": "00b14b89334287ee06c704a79e848aac94780213a91f1528b88e07b17f099540",
}
EFFECTIVE_TEMPERATURE = 0.0
TEMPERATURE_PROVENANCE = (
    "inferred from frozen config.py default STATECLARIFY_TEMPERATURE=0 and "
    "launch-environment audit; legacy runner did not persist this field"
)
DOMAINS = {"env": "environment", "git": "git", "db": "database", "dep": "deployment"}
ARMS = {
    "fin": "statewise_v2_contract_corrected_typedfix",
    "mem": "memlineage_trust",
    "saf": "safety_sentry",
    "ask": "ask_or_assume",
}
LABELS = {
    "statewise_v2_contract_corrected_typedfix": "StateWise",
    "memlineage_trust": "MemLineage",
    "safety_sentry": "Safety Sentry",
    "ask_or_assume": "Ask-or-Assume",
}


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _records(sheet: Any) -> list[dict[str, Any]]:
    values = list(sheet.values)
    if not values:
        return []
    headers = list(values[0])
    return [dict(zip(headers, row)) for row in values[1:]]


def _write_sheet(workbook: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(name)
    if not rows:
        sheet.append(["status"])
        sheet.append(["none"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _expected_sequences(domain: str) -> list[str]:
    return [f"seq_{domain}_{index:03d}" for index in range(1, 9)]


def _validate_manifest(
    path: Path, manifest: dict[str, Any], domain: str, arm: str
) -> None:
    expected = _expected_sequences(domain)
    config = manifest.get("experiment_config") or {}
    errors = []
    if manifest.get("run_id") != path.name:
        errors.append("run id/path mismatch")
    if manifest.get("status") != "COMPLETE":
        errors.append(f"top-level status is {manifest.get('status')!r}, not COMPLETE")
    if config.get("model") != "gpt-5.6-sol":
        errors.append("model mismatch")
    if config.get("fake_api"):
        errors.append("fake API enabled")
    if config.get("sequences") != expected:
        errors.append("sequence membership/order mismatch")
    if config.get("arms") != [arm]:
        errors.append("arm mismatch")
    if config.get("dataset", {}).get("composite_sha256") != DATASET_SHA:
        errors.append("dataset fingerprint mismatch")
    if config.get("measurement_code_fingerprints") != EXPECTED_CODE_HASHES:
        errors.append("measurement code fingerprint mismatch")
    if (manifest.get("manager") or {}).get("max_concurrency") != 1:
        errors.append("worker is not serial")
    episodes = manifest.get("episodes") or {}
    if len(episodes) != 8:
        errors.append(f"expected 8 episodes, found {len(episodes)}")
    statuses = [episode.get("status") for episode in episodes.values()]
    if statuses.count("SUCCEEDED") != 8:
        errors.append(f"expected 8 SUCCEEDED episodes, found {statuses.count('SUCCEEDED')}")
    if errors:
        raise RuntimeError(f"{path.name}: " + "; ".join(errors))


def _load_replication() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    sequence_rows: list[dict[str, Any]] = []
    task_rows: list[dict[str, Any]] = []
    provenance: list[dict[str, Any]] = []
    for domain_code, domain in DOMAINS.items():
        for arm_code, arm in ARMS.items():
            path = RUN_ROOT / f"r5v10-rep3-{domain_code}-{arm_code}"
            manifest_path = path / "manifest.json"
            if not manifest_path.is_file():
                raise RuntimeError(f"missing manifest: {manifest_path}")
            manifest = _read_json(manifest_path)
            _validate_manifest(path, manifest, domain, arm)
            config = manifest["experiment_config"]
            usage = (manifest.get("aggregate") or {}).get("usage") or {}
            provenance.append({
                "run_id": manifest["run_id"],
                "domain": domain,
                "arm": arm,
                "status": manifest["status"],
                "model": config["model"],
                "temperature": EFFECTIVE_TEMPERATURE,
                "temperature_provenance": TEMPERATURE_PROVENANCE,
                "dataset_sha256": config["dataset"]["composite_sha256"],
                "measurement_code_sha256": hashlib.sha256(
                    json.dumps(EXPECTED_CODE_HASHES, sort_keys=True).encode("utf-8")
                ).hexdigest(),
                "manifest_sha256": _sha(manifest_path),
                "calls": int(usage.get("calls") or 0),
                "prompt_tokens": int(usage.get("prompt_tokens") or 0),
                "completion_tokens": int(usage.get("completion_tokens") or 0),
                "total_tokens": int(usage.get("total_tokens") or 0),
                "latency_seconds": float(usage.get("latency_seconds") or 0),
                "manifest": str(manifest_path),
            })
            for episode_id, episode in manifest["episodes"].items():
                result = (episode.get("attempts") or [])[-1].get("result") or {}
                metrics = result.get("metrics") or {}
                sequence_id = result.get("sequence_id")
                if sequence_id not in _expected_sequences(domain):
                    raise RuntimeError(f"{path.name}: unexpected sequence result {sequence_id!r}")
                row = {
                    "run_id": manifest["run_id"],
                    "domain": domain,
                    "arm": arm,
                    "method": LABELS[arm],
                    "sequence_id": sequence_id,
                    "trial": episode["spec"]["trial"],
                    **metrics,
                }
                sequence_rows.append(row)
                for task in result.get("tasks") or []:
                    summary = task.get("summary") or {}
                    task_rows.append({
                        "run_id": manifest["run_id"],
                        "domain": domain,
                        "arm": arm,
                        "method": LABELS[arm],
                        "sequence_id": sequence_id,
                        "task_id": task.get("task_id"),
                        "task_index": task.get("task_index"),
                        "task_success": bool(summary.get("task_success")),
                        "risk_adjusted_task_success": bool(summary.get("risk_adjusted_task_success")),
                        "safety_violation": bool(summary.get("safety_violation")),
                        "repair_performed": summary.get("repair_performed"),
                        "repair_correct": summary.get("repair_correct"),
                        "terminal_reason": summary.get("terminal_reason") or summary.get("status"),
                        "handoff": task.get("inherited_previous_final_exactly"),
                        "state_reuse": task.get("repaired_memory_entered_planner_context"),
                    })
    sequence_keys = [(row["arm"], row["sequence_id"], row["trial"]) for row in sequence_rows]
    task_keys = [(row["arm"], row["sequence_id"], row["task_id"]) for row in task_rows]
    if len(sequence_rows) != 128 or len(sequence_keys) != len(set(sequence_keys)):
        raise RuntimeError("replication sequence matrix is incomplete or duplicated")
    if len(task_rows) != 512 or len(task_keys) != len(set(task_keys)):
        raise RuntimeError("replication task matrix is incomplete or duplicated")
    return sequence_rows, task_rows, provenance


def _summaries(
    sequences: list[dict[str, Any]], tasks: list[dict[str, Any]], provenance: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    rows = []
    for arm in ARMS.values():
        arm_sequences = [row for row in sequences if row["arm"] == arm]
        arm_tasks = [row for row in tasks if row["arm"] == arm]
        future = [row for row in arm_tasks if int(row["task_index"]) > 1]
        sources = [row for row in provenance if row["arm"] == arm]
        rows.append({
            "method": LABELS[arm],
            "arm": arm,
            "sequences": len(arm_sequences),
            "tasks": len(arm_tasks),
            "all_future_safe_completion": sum(
                int(row.get("future_task_success_count") or 0) == 3
                and int(row.get("cumulative_safety_violations") or 0) == 0
                for row in arm_sequences
            ) / len(arm_sequences),
            "future_task_success_rate": sum(bool(row["task_success"]) for row in future) / len(future),
            "future_safe_completion_rate": sum(bool(row["risk_adjusted_task_success"]) for row in future) / len(future),
            "future_no_violation_rate": sum(not row["safety_violation"] for row in future) / len(future),
            "safety_violations": sum(bool(row["safety_violation"]) for row in arm_tasks),
            "state_handoff_integrity_rate": mean(float(row["state_handoff_integrity_rate"]) for row in arm_sequences),
            "state_reuse_rate": mean(float(row["state_reuse_rate"]) for row in arm_sequences),
            "repeated_intervention_rate": mean(float(row["repeated_intervention_rate"]) for row in arm_sequences),
            "api_calls_per_task": sum(row["calls"] for row in sources) / len(arm_tasks),
            "tokens_per_task": sum(row["total_tokens"] for row in sources) / len(arm_tasks),
            "api_latency_seconds_per_task": sum(row["latency_seconds"] for row in sources) / len(arm_tasks),
        })
    return rows


def _agreement(replication: list[dict[str, Any]]) -> list[dict[str, Any]]:
    workbook = load_workbook(PRIMARY, read_only=True, data_only=True)
    primary = [
        row for row in _records(workbook["Sequence_Data"])
        if row.get("arm") in set(ARMS.values())
    ]
    primary_by_key = {(row["arm"], row["sequence_id"]): row for row in primary}
    replication_by_key = {(row["arm"], row["sequence_id"]): row for row in replication}
    if set(primary_by_key) != set(replication_by_key):
        missing = sorted(set(primary_by_key) - set(replication_by_key))
        extra = sorted(set(replication_by_key) - set(primary_by_key))
        raise RuntimeError(f"primary/replication sequence keys differ: missing={missing}, extra={extra}")
    metrics = {
        "all_future_safe_completion": lambda row: (
            int(row.get("future_task_success_count") or 0) == 3
            and int(row.get("cumulative_safety_violations") or 0) == 0
        ),
        "all_future_tasks_succeeded": lambda row: int(
            row.get("future_task_success_count") or 0
        ) == 3,
        "future_no_violation": lambda row: int(row.get("cumulative_safety_violations") or 0) == 0,
        "t1_repaired_contract_satisfied": lambda row: bool(row.get("t1_repaired_contract_satisfied")),
    }
    results = []
    for arm in ARMS.values():
        keys = sorted(key for key in primary_by_key if key[0] == arm)
        for metric, value in metrics.items():
            pairs = [(value(primary_by_key[key]), value(replication_by_key[key])) for key in keys]
            results.append({
                "method": LABELS[arm],
                "arm": arm,
                "metric": metric,
                "paired_sequences": len(pairs),
                "primary_rate": sum(a for a, _ in pairs) / len(pairs),
                "replication_rate": sum(b for _, b in pairs) / len(pairs),
                "absolute_rate_difference": abs(sum(a for a, _ in pairs) - sum(b for _, b in pairs)) / len(pairs),
                "exact_agreement_rate": sum(a == b for a, b in pairs) / len(pairs),
                "primary_only_positive": sum(a and not b for a, b in pairs),
                "replication_only_positive": sum(b and not a for a, b in pairs),
            })
    return results


def main() -> None:
    sequences, tasks, provenance = _load_replication()
    summaries = _summaries(sequences, tasks, provenance)
    agreement = _agreement(sequences)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Replication_Summary", summaries)
    _write_sheet(workbook, "Run_Agreement", agreement)
    _write_sheet(workbook, "Sequence_Data", sequences)
    _write_sheet(workbook, "Task_Data", tasks)
    _write_sheet(workbook, "Run_Provenance", provenance)
    failures = [row for row in tasks if not row["task_success"] or row["safety_violation"]]
    _write_sheet(workbook, "Behavioral_Failures", failures)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    validation = {
        "status": "VALID",
        "measurement_only": True,
        "validated_at": datetime.now(timezone.utc).isoformat(),
        "dataset_sha256": DATASET_SHA,
        "model": "gpt-5.6-sol",
        "temperature": EFFECTIVE_TEMPERATURE,
        "temperature_provenance": TEMPERATURE_PROVENANCE,
        "headline_definition": (
            "future_task_success_count == 3 and cumulative_safety_violations == 0"
        ),
        "worker_isolation": "16 OS processes, one serial shard per process",
        "runs": 16,
        "sequences": 128,
        "tasks": 512,
        "replication_workbook": str(OUTPUT),
        "replication_workbook_sha256": _sha(OUTPUT),
        "primary_workbook": str(PRIMARY),
        "primary_workbook_sha256": _sha(PRIMARY),
        "primary_overwritten": False,
        "summary": summaries,
        "agreement": agreement,
        "source_manifests": {row["run_id"]: row["manifest_sha256"] for row in provenance},
    }
    VALIDATION.write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(OUTPUT.resolve())
    print(VALIDATION.resolve())


if __name__ == "__main__":
    main()
