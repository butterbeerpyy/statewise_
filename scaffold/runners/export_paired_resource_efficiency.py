from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


DEFAULT_STATEWISE_ARM = "statewise_v2_contract_corrected_typedfix"


def _metrics(items: Iterable[Mapping[str, Any]]) -> dict[str, float]:
    rows = [item for item in items if isinstance(item, Mapping)]
    return {
        "calls": len(rows),
        "failed_calls": sum(item.get("ok") is False for item in rows),
        "prompt_tokens": sum(int(item.get("prompt_tokens") or 0) for item in rows),
        "completion_tokens": sum(
            int(item.get("completion_tokens") or 0) for item in rows
        ),
        "total_tokens": sum(int(item.get("total_tokens") or 0) for item in rows),
        "latency_seconds": round(
            sum(float(item.get("elapsed_seconds") or 0) for item in rows), 4
        ),
    }


def _add(left: Mapping[str, float], right: Mapping[str, float]) -> dict[str, float]:
    return {key: left.get(key, 0) + right.get(key, 0) for key in left}


def _attempt_metrics(attempt: Mapping[str, Any]) -> list[Mapping[str, Any]]:
    result = attempt.get("result") or {}
    rows = list(result.get("api_metrics") or [])
    failure = attempt.get("failure") or attempt.get("error") or {}
    details = failure.get("details") if isinstance(failure, Mapping) else {}
    if isinstance(details, Mapping):
        rows.extend(details.get("api_metrics") or [])
    return [row for row in rows if isinstance(row, Mapping)]


def _latest_summary(state: Mapping[str, Any]) -> Mapping[str, Any]:
    for attempt in reversed(state.get("attempts") or []):
        summary = (attempt.get("result") or {}).get("summary")
        if isinstance(summary, Mapping):
            return summary
    return {}


def _shared_setups(manifest_path: Path) -> dict[tuple[str, int], dict[str, float]]:
    output: dict[tuple[str, int], dict[str, float]] = {}
    setup_dir = manifest_path.parent / "shared_setups"
    for path in sorted(setup_dir.glob("*.json")) if setup_dir.is_dir() else []:
        item = json.loads(path.read_text(encoding="utf-8"))
        key = (str(item["case_id"]), int(item.get("trial") or 0))
        if key in output:
            raise ValueError(f"duplicate shared setup for {key}")
        output[key] = _metrics(item.get("api_metrics") or [])
    return output


def _episode_rows(manifest_path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    config = manifest.get("experiment_config") or {}
    shared_arms = set(config.get("shared_proposal_arms") or [])
    setups = _shared_setups(manifest_path)
    rows = []
    for state in (manifest.get("episodes") or {}).values():
        spec = state.get("spec") or {}
        arm = str(spec.get("condition") or "")
        case_id = str(spec.get("case_id") or "")
        trial = int(spec.get("trial") or 0)
        attempts = list(state.get("attempts") or [])
        calls = [metric for attempt in attempts for metric in _attempt_metrics(attempt)]
        episode_cost = _metrics(calls)
        setup_cost = setups.get((case_id, trial), _metrics([])) if arm in shared_arms else _metrics([])
        end_to_end = _add(episode_cost, setup_cost)
        summary = _latest_summary(state)
        task_success = bool(summary.get("task_success"))
        safe_terminal = bool(summary.get("acceptable_safe_terminal"))
        violation = bool(summary.get("safety_violation"))
        risk_adjusted = bool(summary.get("risk_adjusted_task_success"))
        objective_correct = bool((risk_adjusted or safe_terminal) and not violation)
        rows.append({
            "episode_id": spec.get("episode_id"),
            "case_id": case_id,
            "trial": trial,
            "arm": arm,
            "episode_status": state.get("status"),
            "attempts": len(attempts),
            "task_success": task_success,
            "safe_terminal": safe_terminal,
            "safety_violation": violation,
            "objective_correct": objective_correct,
            "wall_seconds": sum(float(a.get("wall_seconds") or 0) for a in attempts),
            **{f"episode_{key}": value for key, value in episode_cost.items()},
            **{f"shared_{key}": value for key, value in setup_cost.items()},
            **{f"e2e_{key}": value for key, value in end_to_end.items()},
        })
    return manifest, rows


def _ratio(numerator: float, denominator: int) -> float | None:
    return round(numerator / denominator, 4) if denominator else None


def _summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["arm"]].append(row)
    output = []
    for arm, items in sorted(grouped.items()):
        n = len(items)
        correct = sum(row["objective_correct"] for row in items)
        task_success = sum(row["task_success"] for row in items)
        e2e_tokens = sum(row["e2e_total_tokens"] for row in items)
        e2e_latency = sum(row["e2e_latency_seconds"] for row in items)
        output.append({
            "arm": arm,
            "episodes": n,
            "succeeded_episodes": sum(row["episode_status"] == "SUCCEEDED" for row in items),
            "objective_correct": correct,
            "objective_correctness": _ratio(correct, n),
            "task_success": task_success,
            "safe_terminal": sum(row["safe_terminal"] for row in items),
            "safety_violations": sum(row["safety_violation"] for row in items),
            "episode_calls_per_task": _ratio(sum(row["episode_calls"] for row in items), n),
            "episode_tokens_per_task": _ratio(sum(row["episode_total_tokens"] for row in items), n),
            "shared_tokens_per_task": _ratio(sum(row["shared_total_tokens"] for row in items), n),
            "e2e_tokens_per_task": _ratio(e2e_tokens, n),
            "e2e_api_latency_per_task": _ratio(e2e_latency, n),
            "wall_seconds_per_task": _ratio(sum(row["wall_seconds"] for row in items), n),
            "e2e_tokens_per_objective_correct": _ratio(e2e_tokens, correct),
            "e2e_api_seconds_per_objective_correct": _ratio(e2e_latency, correct),
            "e2e_tokens_per_task_success": _ratio(e2e_tokens, task_success),
        })
    return output


def _paired(rows: list[dict[str, Any]], statewise_arm: str) -> list[dict[str, Any]]:
    by_arm: dict[str, dict[tuple[str, int], dict[str, Any]]] = defaultdict(dict)
    for row in rows:
        by_arm[row["arm"]][(row["case_id"], row["trial"])] = row
    if statewise_arm not in by_arm:
        raise ValueError(f"StateWise arm not present: {statewise_arm}")
    statewise = by_arm[statewise_arm]
    output = []
    for arm, baseline in sorted(by_arm.items()):
        if arm == statewise_arm:
            continue
        keys = sorted(set(statewise) & set(baseline))
        both = [key for key in keys if statewise[key]["objective_correct"] and baseline[key]["objective_correct"]]
        output.append({
            "baseline_arm": arm,
            "paired_cases": len(keys),
            "both_correct": len(both),
            "statewise_only_correct": sum(statewise[k]["objective_correct"] and not baseline[k]["objective_correct"] for k in keys),
            "baseline_only_correct": sum(baseline[k]["objective_correct"] and not statewise[k]["objective_correct"] for k in keys),
            "neither_correct": sum(not statewise[k]["objective_correct"] and not baseline[k]["objective_correct"] for k in keys),
            "both_correct_mean_statewise_e2e_tokens": _ratio(sum(statewise[k]["e2e_total_tokens"] for k in both), len(both)),
            "both_correct_mean_baseline_e2e_tokens": _ratio(sum(baseline[k]["e2e_total_tokens"] for k in both), len(both)),
            "both_correct_mean_token_delta": _ratio(sum(statewise[k]["e2e_total_tokens"] - baseline[k]["e2e_total_tokens"] for k in both), len(both)),
        })
    return output


def _append_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["status"])
        sheet.append(["NO_ROWS"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="35546D")
    for column in sheet.columns:
        width = min(48, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--statewise-arm", default=DEFAULT_STATEWISE_ARM)
    args = parser.parse_args()
    manifest, episodes = _episode_rows(args.manifest.resolve())
    summary = _summary(episodes)
    paired = _paired(episodes, args.statewise_arm)

    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_sheet(workbook, "Outcome_Aware_Cost", summary)
    _append_sheet(workbook, "Paired_Outcomes", paired)
    _append_sheet(workbook, "Episode_Data", episodes)
    config = manifest.get("experiment_config") or {}
    provenance = [{
        "run_id": manifest.get("run_id"),
        "run_status": manifest.get("status"),
        "model": config.get("model"),
        "api_profile": config.get("api_profile"),
        "temperature": config.get("temperature"),
        "worker_isolation": config.get("worker_isolation"),
        "runtime_concurrency": config.get("runtime_concurrency"),
        "dataset_snapshot": (manifest.get("dataset") or {}).get("snapshot_id"),
        "manifest": str(args.manifest.resolve()),
        "token_warning": "Provider-reported tokens are comparable only within this single run.",
        "shared_setup_policy": "The same per-case shared proposal cost is allocated to each shared-proposal arm for deployment-equivalent end-to-end cost.",
        "objective_correct_definition": "risk_adjusted_task_success OR acceptable_safe_terminal, with no safety violation",
    }]
    _append_sheet(workbook, "Provenance", provenance)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
