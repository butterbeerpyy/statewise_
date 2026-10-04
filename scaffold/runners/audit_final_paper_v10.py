from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
PAPER = ROOT / "paper/fse2027"
RESULTS = ROOT / 'paper_experiments/results/shared'
MATCHED = RESULTS / "permission_matched_v10_all_arms_20260819.xlsx"
MATCHED_AUDIT = RESULTS / "audit/PERMISSION_MATCHED_V10_ALL_ARMS_20260819_AUDIT.json"
COMPUTE_AUDIT = RESULTS / "audit/COMPUTE_MATCHED_REPETITION_V10_20260819_AUDIT.json"
OUTPUT = RESULTS / "audit/FINAL_PAPER_V10_AUDIT.json"

LABELS = {
    "correct": "Correct memory",
    "corrupted": "Corrupted memory",
    "no_memory": "No memory",
    "bare_approval": "Bare approval",
    "rationale_approval": "Rationale approval",
    "generic_clarification": "Generic clarification",
    "memlineage_trust": "MemLineage",
    "safety_sentry": "Safety Sentry",
    "ask_or_assume": "Ask-or-Assume",
    "statewise_v2_contract_corrected_typedfix": "\\statewise{}",
}


def records(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    sheet = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(values)]
    return [dict(zip(headers, row)) for row in values]


def table_block(text: str, label: str) -> str:
    label_position = text.index(f"\\label{{{label}}}")
    start = text.rfind("\\begin{table}", 0, label_position)
    end = text.index("\\end{table}", label_position) + len("\\end{table}")
    return text[start:end]


def pct(value: Any) -> str:
    return f"{100 * float(value):.1f}\\%"


def main() -> int:
    evaluation = (PAPER / "sections/evaluation.tex").read_text(encoding="utf-8")
    main_tex = (PAPER / "main.tex").read_text(encoding="utf-8")
    discussion = (PAPER / "sections/discussion.tex").read_text(encoding="utf-8")
    appendix = (PAPER / "sections/appendix.tex").read_text(encoding="utf-8")
    limitations = (PAPER / "sections/limitations.tex").read_text(encoding="utf-8")
    paper_text = "\n".join(
        path.read_text(encoding="utf-8")
        for path in [PAPER / "main.tex", *sorted((PAPER / "sections").glob("*.tex"))]
    )

    matched_audit = json.loads(MATCHED_AUDIT.read_text(encoding="utf-8"))
    compute_audit = (
        json.loads(COMPUTE_AUDIT.read_text(encoding="utf-8"))
        if COMPUTE_AUDIT.exists()
        else {"status": "PENDING"}
    )
    summary = records(MATCHED, "Paper_Summary")
    domains = records(MATCHED, "Domain_Summary")
    main_table = table_block(evaluation, "tab:main")
    domain_table = table_block(evaluation, "tab:domains")

    main_rows_match = all(
        LABELS[row["arm"]] in main_table
        and pct(row["objective_correctness"]) in main_table
        and pct(row["action_completion"]) in main_table
        and pct(row["safety_violation"]) in main_table
        for row in summary
    )
    displayed_domain_arms = {
        "correct",
        "corrupted",
        "no_memory",
        "memlineage_trust",
        "safety_sentry",
        "ask_or_assume",
        "statewise_v2_contract_corrected_typedfix",
    }
    domain_rows_match = all(
        pct(row["behavioral_correctness_rate"]) in domain_table
        and pct(row["safety_violation_rate"]) in domain_table
        for row in domains
        if row["arm"] in displayed_domain_arms
    )

    checks = {
        "permission_matched_audit_valid": matched_audit.get("status") == "VALID",
        "permission_replacement_recommended": matched_audit.get("replacement_recommended") is True,
        "statewise_zero_violations": matched_audit.get("checks", {}).get("statewise_zero_violations") is True,
        "statewise_best_deployable": matched_audit.get("checks", {}).get("statewise_best_deployable") is True,
        "main_table_matches_workbook": main_rows_match,
        "domain_table_matches_workbook": domain_rows_match,
        "no_count_fractions_in_paper": re.search(r"(?<![A-Za-z0-9])\d+\s*/\s*\d+(?![A-Za-z0-9])", paper_text) is None,
        "no_n_table_header": not re.search(r"(?:&|\{)\s*[nN]\s*(?:&|\\\\|\})", paper_text),
        "headline_updated": all(
            token in main_tex + evaluation
            for token in ("95.3\\%", "94.4\\%", "80.7\\%", "38.7\\%")
        ),
        "stale_headlines_absent": all(
            token not in main_tex + evaluation
            for token in ("69.3\\%", "26.7\\%", "recovers correctness to 84.0\\%")
        ),
        "correct_memory_failure_breakdown_updated": all(
            token in discussion for token in ("19.3\\%", "62.1\\%", "31.0\\%", "6.9\\%")
        ),
        "matched_permission_disclosed": all(
            token in evaluation + limitations + appendix
            for token in ("all ten", "post-freeze", "permission")
        ),
        "true_residual_gate_violation_disclosed": all(
            token in evaluation for token in ("recursively delete", "protected directory", "one event")
        ),
        "compute_matched_audit_valid": compute_audit.get("status") == "VALID",
        "compute_matched_results_reported": all(
            token in evaluation + discussion + appendix
            for token in ("79.1\\%", "38.0\\%", "84.0\\%", "45.3\\%", "107.2\\%", "100.7\\%")
        ),
        "compute_violation_disclosed": all(
            token in appendix for token in ("already-present dependency", "protected package file", "real")
        ),
        "review_build_anonymous": "Anonymous Author(s)" in (PAPER / "fse-review.tex").read_text(encoding="utf-8")
        and 'anonymous' not in (PAPER / "fse-review.tex").read_text(encoding="utf-8"),
        "arxiv_build_identified": all(
            token in (PAPER / "arxiv.tex").read_text(encoding="utf-8")
            for token in ('personal-author-name', 'personal-institution', 'personal-email-address')
        ),
    }
    payload = {
        "status": "VALID" if all(checks.values()) else "INVALID",
        "checks": checks,
        "permission_audit": str(MATCHED_AUDIT.relative_to(ROOT)),
        "compute_audit": str(COMPUTE_AUDIT.relative_to(ROOT)),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if payload["status"] == "VALID" else 1


if __name__ == "__main__":
    raise SystemExit(main())
