from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


EXPECTED_RUN_ID = "permission-sensitivity-v10-statewise20-gpt56sol-20260815-r1"
EXPECTED_ARM = "statewise_v2_contract_corrected_typedfix"
EXPECTED_DATASET_ID = "w1_contract_corrected_v10_permission_enabled_sensitivity"
EXPECTED_DATASET_SHA = "eefddd2ec3643136bf6bba0ecb395275e690212f55a07c028a39b8e789daed79"
EXPECTED_METHOD_FP = "sha256:0804b97693cb2f4cab236f325347b3c227c391572aa7e3048f85657640cc899e"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _rows_from_sheet(path: Path, sheet: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    values = worksheet.iter_rows(values_only=True)
    header = next(values)
    return [dict(zip(header, row)) for row in values]


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title)
    if not rows:
        worksheet.append(["no_rows"])
        return
    fields = list(rows[0])
    worksheet.append(fields)
    for row in rows:
        worksheet.append([row.get(field) for field in fields])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def finalize(
    run_dir: Path,
    overlay_snapshot: Path,
    formal_xlsx: Path,
    output_xlsx: Path,
    output_audit: Path,
    expected_run_id: str = EXPECTED_RUN_ID,
    comparison_manifest: Path | None = None,
) -> dict[str, Any]:
    manifest_path = run_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    overlay = json.loads(overlay_snapshot.read_text(encoding="utf-8"))
    cohort = set(overlay["sensitivity_cohort_case_ids"])
    formal_rows = _rows_from_sheet(formal_xlsx, "Episode_Data")
    original = {
        row["case_id"]: row
        for row in formal_rows
        if row.get("arm") == EXPECTED_ARM and row.get("case_id") in cohort
    }

    records = []
    episode_statuses = Counter()
    api_failed_calls = 0
    for state in manifest["episodes"].values():
        case_id = state["spec"]["case_id"]
        episode_statuses[state["status"]] += 1
        attempt = state.get("attempts", [])[-1] if state.get("attempts") else {}
        result = attempt.get("result", {})
        summary = result.get("summary", {})
        api_metrics = result.get("api_metrics", [])
        failed = sum(item.get("ok") is False for item in api_metrics)
        api_failed_calls += failed
        before = original.get(case_id, {})
        records.append({
            "case_id": case_id,
            "episode_status": state["status"],
            "attempt_status": attempt.get("status"),
            "formal_v9_objective_correct": bool(before.get("behaviorally_correct")),
            "formal_v9_action_completed": bool(before.get("action_completed")),
            "permission_enabled_task_success": bool(summary.get("task_success")),
            "permission_enabled_safety_violation": bool(summary.get("safety_violation")),
            "repair_performed": bool(summary.get("repair_performed")),
            "repair_correct": summary.get("repair_correct"),
            "repair_event_count": summary.get("repair_event_count"),
            "permission_interactions": summary.get("permission_interactions"),
            "final_operation": (summary.get("final_action") or {}).get("operation"),
            "api_calls": len(api_metrics),
            "api_failed_calls": failed,
            "changed_from_formal_v9": bool(summary.get("task_success"))
            != bool(before.get("action_completed")),
        })
    records.sort(key=lambda row: row["case_id"])

    config = manifest.get("experiment_config", {})
    dataset = manifest.get("dataset", {})
    implementation = (
        config.get("arm_registry", {})
        .get(EXPECTED_ARM, {})
        .get("implementation", {})
    )
    selected = set(config.get("selected_cases", []))
    checks = {
        "run_id_matches": manifest.get("run_id") == expected_run_id,
        "run_complete": manifest.get("status") == "COMPLETE",
        "exactly_20_succeeded": episode_statuses == Counter({"SUCCEEDED": 20}),
        "selected_cases_match_cohort": selected == cohort and len(cohort) == 20,
        "formal_v9_pairing_complete": set(original) == cohort,
        "dataset_id_matches": dataset.get("snapshot_id") == EXPECTED_DATASET_ID,
        "dataset_sha_matches": dataset.get("snapshot_sha256") == EXPECTED_DATASET_SHA,
        "overlay_file_sha_matches": _sha256(overlay_snapshot) == EXPECTED_DATASET_SHA,
        "method_fingerprint_matches": implementation.get("fingerprint")
        == EXPECTED_METHOD_FP,
        "process_isolated": config.get("worker_isolation") == "process",
        "single_arm": config.get("arms") == [EXPECTED_ARM],
        "live_gpt56sol": config.get("model") == "gpt-5.6-sol"
        and config.get("fake_api") is False,
        "temperature_zero": float(config.get("temperature")) == 0.0,
        "no_api_failures": api_failed_calls == 0,
    }
    status = "VALID" if all(checks.values()) else "INVALID"
    completed = sum(row["permission_enabled_task_success"] for row in records)
    violations = sum(row["permission_enabled_safety_violation"] for row in records)
    improved = sum(
        row["permission_enabled_task_success"]
        and not row["formal_v9_action_completed"]
        for row in records
    )
    summary = [{
        "reporting_status": "sensitivity_only_not_primary_heldout",
        "n": len(records),
        "formal_v9_action_completed_n": sum(
            row["formal_v9_action_completed"] for row in records
        ),
        "formal_v9_action_completion": (
            sum(row["formal_v9_action_completed"] for row in records) / len(records)
            if records else None
        ),
        "permission_enabled_completed_n": completed,
        "permission_enabled_completion": completed / len(records) if records else None,
        "paired_improved_n": improved,
        "safety_violation_n": violations,
        "safety_violation_rate": violations / len(records) if records else None,
        "interpretation": (
            "Construct-validity sensitivity: isolates explicit deterministic "
            "permission availability; does not replace the frozen 150-case table."
        ),
    }]
    provenance = [{
        "run_id": manifest.get("run_id"),
        "run_status": manifest.get("status"),
        "dataset_snapshot_id": dataset.get("snapshot_id"),
        "dataset_snapshot_sha256": dataset.get("snapshot_sha256"),
        "parent_snapshot_sha256": overlay.get("parent_snapshot", {}).get("sha256"),
        "method_fingerprint": implementation.get("fingerprint"),
        "model": config.get("model"),
        "temperature": config.get("temperature"),
        "worker_isolation": config.get("worker_isolation"),
        "runtime_concurrency": config.get("runtime_concurrency"),
        "api_profile": config.get("api_profile"),
        "json_response_format": config.get("json_response_format"),
        "validation_status": status,
    }]

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Sensitivity_Summary", summary)
    _write_sheet(workbook, "Paired_Case_Data", records)
    _write_sheet(workbook, "Validation_Checks", [
        {"check": key, "passed": value} for key, value in checks.items()
    ])
    _write_sheet(workbook, "Provenance", provenance)
    workbook.save(output_xlsx)

    audit = {
        "schema_version": "permission-enabled-sensitivity-result-audit-v1",
        "status": status,
        "reporting_status": "sensitivity_only_not_primary_heldout",
        "checks": checks,
        "episode_statuses": dict(episode_statuses),
        "summary": summary[0],
        "provenance": provenance[0],
        "transport": {
            "json_response_format": config.get("json_response_format"),
            "json_only_prompt_and_parser_retained": True,
        },
        "artifacts": {
            "manifest": str(manifest_path),
            "overlay_snapshot": str(overlay_snapshot),
            "formal_v9_xlsx": str(formal_xlsx),
            "output_xlsx": str(output_xlsx),
        },
    }
    if comparison_manifest is not None:
        comparison = json.loads(comparison_manifest.read_text(encoding="utf-8"))
        prior = {}
        for state in comparison["episodes"].values():
            attempt = state.get("attempts", [])[-1] if state.get("attempts") else {}
            prior_summary = attempt.get("result", {}).get("summary", {})
            prior[state["spec"]["case_id"]] = {
                "task_success": bool(prior_summary.get("task_success")),
                "safety_violation": bool(prior_summary.get("safety_violation")),
            }
        current = {
            row["case_id"]: {
                "task_success": row["permission_enabled_task_success"],
                "safety_violation": row["permission_enabled_safety_violation"],
            }
            for row in records
        }
        common = sorted(set(prior) & set(current))
        agreement = sum(prior[case_id] == current[case_id] for case_id in common)
        audit["replication_comparison"] = {
            "comparison_run_id": comparison.get("run_id"),
            "same_case_ids": set(prior) == set(current),
            "compared_case_count": len(common),
            "case_level_outcome_agreement_count": agreement,
            "case_level_outcome_agreement_rate": (
                agreement / len(common) if common else None
            ),
            "aggregate_completion_matches": (
                sum(row["task_success"] for row in prior.values())
                == sum(row["task_success"] for row in current.values())
            ),
            "aggregate_safety_matches": (
                sum(row["safety_violation"] for row in prior.values())
                == sum(row["safety_violation"] for row in current.values())
            ),
        }
        audit["artifacts"]["comparison_manifest"] = str(comparison_manifest)
    output_audit.parent.mkdir(parents=True, exist_ok=True)
    output_audit.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return audit


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--overlay-snapshot", type=Path, required=True)
    parser.add_argument("--formal-xlsx", type=Path, required=True)
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-audit", type=Path, required=True)
    parser.add_argument("--expected-run-id", default=EXPECTED_RUN_ID)
    parser.add_argument("--comparison-manifest", type=Path)
    args = parser.parse_args()
    audit = finalize(
        args.run_dir.resolve(),
        args.overlay_snapshot.resolve(),
        args.formal_xlsx.resolve(),
        args.output_xlsx.resolve(),
        args.output_audit.resolve(),
        args.expected_run_id,
        args.comparison_manifest.resolve() if args.comparison_manifest else None,
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0 if audit["status"] == "VALID" else 1


if __name__ == "__main__":
    raise SystemExit(main())
