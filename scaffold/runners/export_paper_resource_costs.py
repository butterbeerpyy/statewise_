from __future__ import annotations

from pathlib import Path
from math import floor, ceil
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / 'paper_experiments/results'
OUTPUT = RESULTS / "shared/rq1_rq5_resource_costs.xlsx"

LABELS = {
    "statewise_v2_contract_corrected_typedfix": "StateWise",
    "memlineage_trust": "MemLineage",
    "safety_sentry": "Safety Sentry",
    "ask_or_assume": "Ask-or-Assume",
    "correct": "Correct memory",
    "corrupted": "Corrupted memory",
    "no_memory": "No memory",
}
PAPER_METHODS = {"StateWise", "MemLineage", "Safety Sentry", "Ask-or-Assume"}


def records(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    sheet = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    values = list(sheet.values)
    headers = list(values[0])
    return [dict(zip(headers, row)) for row in values[1:]]


def percentile(values: list[float], q: float) -> float | None:
    values = sorted(float(value) for value in values if value is not None)
    if not values:
        return None
    position = (len(values) - 1) * q
    lower, upper = floor(position), ceil(position)
    if lower == upper:
        return values[lower]
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def episode_distributions(path: Path, sheet_name: str) -> dict[str, dict[str, float | None]]:
    rows = records(path, sheet_name)
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        status = str(row.get("episode_status") or row.get("status") or "").upper()
        if status and status != "SUCCEEDED":
            continue
        grouped.setdefault(str(row.get("arm") or ""), []).append(row)

    output = {}
    for arm, arm_rows in grouped.items():
        token_key = "total_tokens" if "total_tokens" in arm_rows[0] else "api_total_tokens"
        latency_key = "api_latency_seconds"
        failed_key = "failed_api_calls" if "failed_api_calls" in arm_rows[0] else "api_failed_calls"
        failed_calls = [float(row.get(failed_key) or 0) for row in arm_rows]
        output[arm] = {
            "mean_failed_api_calls": sum(failed_calls) / len(failed_calls),
            "p95_failed_api_calls": percentile(failed_calls, 0.95),
            "p50_total_tokens": percentile([row.get(token_key) for row in arm_rows], 0.50),
            "p95_total_tokens": percentile([row.get(token_key) for row in arm_rows], 0.95),
            "p50_api_latency_seconds": percentile([row.get(latency_key) for row in arm_rows], 0.50),
            "p95_api_latency_seconds": percentile([row.get(latency_key) for row in arm_rows], 0.95),
            "p50_wall_seconds": percentile([row.get("wall_seconds") for row in arm_rows], 0.50),
            "p95_wall_seconds": percentile([row.get("wall_seconds") for row in arm_rows], 0.95),
        }
    return output


def normalized_row(
    rq: str,
    row: dict[str, Any],
    source: Path,
    evidence_run: str,
    distribution: dict[str, float | None] | None = None,
) -> dict[str, Any]:
    arm = str(row.get("arm") or "")
    label = str(row.get("method") or LABELS.get(arm, arm))
    episodes = row.get(
        "episodes_succeeded",
        row.get("episodes", row.get("episodes_total", row.get("tasks"))),
    )
    calls = row.get("mean_api_calls", row.get("api_calls_per_task"))
    prompt = row.get("mean_prompt_tokens")
    completion = row.get("mean_completion_tokens")
    total = row.get("mean_total_tokens", row.get("mean_tokens", row.get("tokens_per_task")))
    latency = row.get("mean_api_latency_seconds", row.get("api_latency_seconds_per_task"))
    wall = row.get("mean_wall_seconds")
    distribution = distribution or {}
    return {
        "rq": rq,
        "evidence_run": evidence_run,
        "method": label,
        "internal_arm": arm,
        "model": row.get("model", "gpt-5.6-sol"),
        "measurement_unit": "task" if rq == "RQ5" else "episode",
        "observations": episodes,
        "mean_api_calls": calls,
        "mean_failed_api_calls": distribution.get("mean_failed_api_calls"),
        "p95_failed_api_calls": distribution.get("p95_failed_api_calls"),
        "mean_prompt_tokens": prompt,
        "mean_completion_tokens": completion,
        "mean_total_tokens": total,
        "mean_api_latency_seconds": latency,
        "mean_wall_seconds": wall,
        "p50_total_tokens": distribution.get("p50_total_tokens"),
        "p95_total_tokens": distribution.get("p95_total_tokens"),
        "p50_api_latency_seconds": distribution.get("p50_api_latency_seconds"),
        "p95_api_latency_seconds": distribution.get("p95_api_latency_seconds"),
        "p50_wall_seconds": distribution.get("p50_wall_seconds"),
        "p95_wall_seconds": distribution.get("p95_wall_seconds"),
        "paper_facing_method": label in PAPER_METHODS,
        "source_workbook": str(source.relative_to(ROOT)),
    }


def main() -> int:
    sources = [
        ("RQ1", RESULTS / "rq1/rq1_results.xlsx", "Arm_Summary", "primary", "Episode_Data"),
        ("RQ2", RESULTS / "rq2/rq2_results.xlsx", "Arm_Summary", "primary", "Episode_Data"),
        ("RQ3", RESULTS / "rq3/rq3_results.xlsx", "Arm_Summary", "primary", "Episode_Data"),
        ("RQ4", RESULTS / "rq4/rq4_results.xlsx", "Arm_Summary", "primary", "Episode_Data"),
        ("RQ5", RESULTS / "rq5/rq5_v10_matched.xlsx", "Paper_Comparison", "primary", None),
        ("RQ5", RESULTS / "rq5/rq5_v10_replication3.xlsx", "Replication_Summary", "replication_3", None),
    ]
    rows = []
    for rq, path, sheet, evidence_run, detail_sheet in sources:
        distributions = episode_distributions(path, detail_sheet) if detail_sheet else {}
        rows.extend(
            normalized_row(
                rq, row, path, evidence_run, distributions.get(str(row.get("arm") or ""))
            )
            for row in records(path, sheet)
        )

    workbook = Workbook()
    data = workbook.active
    data.title = "Cost_By_RQ_Arm"
    headers = list(rows[0])
    data.append(headers)
    for row in rows:
        data.append([row.get(header) for header in headers])

    plot = workbook.create_sheet("Plot_Data")
    plot_headers = [
        "rq", "evidence_run", "method", "mean_api_calls", "mean_failed_api_calls",
        "p95_failed_api_calls", "mean_total_tokens",
        "mean_api_latency_seconds", "mean_wall_seconds", "p50_total_tokens",
        "p95_total_tokens", "p50_wall_seconds", "p95_wall_seconds",
        "paper_facing_method",
    ]
    plot.append(plot_headers)
    for row in rows:
        plot.append([row.get(header) for header in plot_headers])

    notes = workbook.create_sheet("Notes")
    notes.append(["field", "value"])
    notes.append(["scope", "Costs are arithmetic means over successful measured tasks/episodes in each source workbook."])
    notes.append(["non_additivity", "Do not sum observations or costs across RQs: canonical RQ workbooks can reuse the same underlying episodes."])
    notes.append(["latency", "API latency is provider-observed latency, not end-to-end wall time."])
    notes.append(["wall_time", "End-to-end episode wall time is available for RQ2/RQ4 only; blank cells mean not measured, not zero."])
    notes.append(["distribution", "P50/P95 use successful episode-level records. RQ5 retained usage only at shard/run aggregate level, so task-level percentiles are unavailable."])
    notes.append(["retry_overhead", "Failed API-call means/P95 are measured from episode-level call metrics for RQ1--RQ4. RQ5 did not retain failed-call counts separately."])
    notes.append(["missing_token_split", "RQ1/RQ3 legacy summaries retain total tokens but not prompt/completion splits."])
    notes.append(["RQ5_provenance", "RQ5 primary and independent replication are separate rows and must not be pooled into an artificial N=64."])

    validation = workbook.create_sheet("Validation")
    validation.append(["check", "passed", "detail"])
    rq2_paper_rows = [
        row for row in rows
        if row["rq"] == "RQ2" and row["paper_facing_method"]
    ]
    validation.append([
        "RQ2 paper-facing rows have wall time",
        all(row["mean_wall_seconds"] is not None for row in rq2_paper_rows),
        f"rows={len(rq2_paper_rows)}",
    ])
    rq5_runs = {
        row["evidence_run"] for row in rows
        if row["rq"] == "RQ5" and row["method"] == "StateWise"
    }
    validation.append([
        "RQ5 primary and replication retained separately",
        rq5_runs == {"primary", "replication_3"},
        ", ".join(sorted(rq5_runs)),
    ])
    rq5_rows = [row for row in rows if row["rq"] == "RQ5"]
    validation.append([
        "RQ5 unavailable task percentiles remain blank",
        all(row["p95_total_tokens"] is None for row in rq5_rows),
        "Shard-level usage cannot recover task-level percentiles.",
    ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="35546D")
        for column in sheet.columns:
            width = min(52, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(OUTPUT.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
