from __future__ import annotations

import json
import math
import random
import shutil
from pathlib import Path
from statistics import mean
from typing import Iterable

from openpyxl import load_workbook
from scipy.stats import binomtest


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / 'paper_experiments/results/rq5/rq5_v10_matched.xlsx'
PUBLIC_WORKBOOK = ROOT / 'paper_experiments/results/rq5/rq5_results.xlsx'
METHOD = "statewise_v2_contract_corrected_typedfix"
BASELINES = ("memlineage_trust", "safety_sentry", "ask_or_assume")


def records(sheet) -> list[dict]:
    values = list(sheet.values)
    headers = list(values[0])
    return [dict(zip(headers, row)) for row in values[1:]]


def future_safe_complete(row: dict) -> bool:
    return math.isclose(float(row["future_task_success_rate"]), 1.0) and math.isclose(
        float(row.get("future_no_violation_rate", row.get("future_task_safety_rate"))), 1.0
    )


def bootstrap_ci(differences: Iterable[int], seed: int = 20260814) -> tuple[float, float]:
    values = list(differences)
    rng = random.Random(seed)
    draws = sorted(
        mean(rng.choice(values) for _ in values)
        for _ in range(20_000)
    )
    return draws[int(0.025 * len(draws))], draws[int(0.975 * len(draws))]


def holm(rows: list[dict]) -> None:
    ordered = sorted(enumerate(rows), key=lambda item: item[1]["mcnemar_p"])
    adjusted = [0.0] * len(rows)
    running = 0.0
    total = len(rows)
    for rank, (index, row) in enumerate(ordered):
        running = max(running, min(1.0, (total - rank) * row["mcnemar_p"]))
        adjusted[index] = running
    for row, value in zip(rows, adjusted):
        row["holm_adjusted_p"] = value


def main() -> int:
    workbook = load_workbook(WORKBOOK)
    rows = records(workbook["Sequence_Data"])
    by_arm = {
        arm: {row["sequence_id"]: row for row in rows if row["arm"] == arm}
        for arm in (METHOD, *BASELINES)
    }
    comparisons = []
    for baseline in BASELINES:
        if set(by_arm[METHOD]) != set(by_arm[baseline]):
            raise ValueError(f"unpaired sequence ids: {baseline}")
        outcomes = [
            (
                future_safe_complete(by_arm[METHOD][sequence_id]),
                future_safe_complete(by_arm[baseline][sequence_id]),
            )
            for sequence_id in sorted(by_arm[METHOD])
        ]
        method_only = sum(method and not control for method, control in outcomes)
        baseline_only = sum(control and not method for method, control in outcomes)
        discordant = method_only + baseline_only
        p_value = (
            binomtest(min(method_only, baseline_only), discordant, 0.5).pvalue
            if discordant else 1.0
        )
        differences = [int(method) - int(control) for method, control in outcomes]
        low, high = bootstrap_ci(differences)
        comparisons.append({
            "method": METHOD,
            "baseline": baseline,
            "sequences": len(outcomes),
            "method_all_future_safe_complete": sum(method for method, _ in outcomes),
            "baseline_all_future_safe_complete": sum(control for _, control in outcomes),
            "paired_risk_difference": mean(differences),
            "bootstrap_ci_low": low,
            "bootstrap_ci_high": high,
            "method_only_success": method_only,
            "baseline_only_success": baseline_only,
            "mcnemar_p": p_value,
        })
    holm(comparisons)

    if "Paired_Sequence_Tests" in workbook.sheetnames:
        del workbook["Paired_Sequence_Tests"]
    sheet = workbook.create_sheet("Paired_Sequence_Tests")
    headers = list(comparisons[0])
    sheet.append(headers)
    for row in comparisons:
        sheet.append([row[name] for name in headers])
    workbook.save(WORKBOOK)
    shutil.copyfile(WORKBOOK, PUBLIC_WORKBOOK)

    output = WORKBOOK.with_name("RQ5_v10_paired.json")
    output.write_text(json.dumps(comparisons, indent=2) + "\n", encoding="utf-8")
    print(output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
