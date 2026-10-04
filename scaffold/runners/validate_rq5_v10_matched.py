from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
RESULT_ROOT = ROOT / 'paper_experiments/results/rq5'
WORKBOOK = RESULT_ROOT / "rq5_v10_matched.xlsx"
PUBLIC_WORKBOOK = RESULT_ROOT / "rq5_results.xlsx"
PAIRED = RESULT_ROOT / "RQ5_v10_paired.json"
REPORT = RESULT_ROOT / "RQ5_V10_MATCHED_VALIDATION.json"

METHODS = {
    "StateWise": "statewise_v2_contract_corrected_typedfix",
    "MemLineage": "memlineage_trust",
    "Safety Sentry": "safety_sentry",
    "Ask-or-Assume": "ask_or_assume",
}
DATASET_SHA = "a871a91427dd4a5bead0115be695bb74df28a660b25990103cdece681b4b7a62"


def records(sheet) -> list[dict]:
    values = list(sheet.values)
    if not values:
        return []
    headers = list(values[0])
    return [dict(zip(headers, row)) for row in values[1:]]


def close(actual, expected) -> bool:
    return math.isclose(float(actual), float(expected), rel_tol=1e-12, abs_tol=1e-12)


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def main() -> int:
    require(
        hashlib.sha256(WORKBOOK.read_bytes()).digest()
        == hashlib.sha256(PUBLIC_WORKBOOK.read_bytes()).digest(),
        "public RQ5 workbook is not the canonical matched export",
    )
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    required_sheets = {
        "Paper_Comparison", "Sequence_Data", "Task_Data", "Run_Provenance",
        "Paired_Sequence_Tests",
    }
    require(required_sheets <= set(workbook.sheetnames), "missing required workbook sheets")

    paper = records(workbook["Paper_Comparison"])
    sequences = records(workbook["Sequence_Data"])
    tasks = records(workbook["Task_Data"])
    provenance = records(workbook["Run_Provenance"])
    paired_sheet = records(workbook["Paired_Sequence_Tests"])
    paired_json = json.loads(PAIRED.read_text(encoding="utf-8"))

    require({row["method"] for row in paper} == set(METHODS), "paper arm set mismatch")
    require(len(paper) == 4, "paper table must contain exactly four methods")
    require(all(row["dataset_sha256"] == DATASET_SHA for row in provenance), "dataset SHA mismatch")
    require(all(row["status"] in {"COMPLETE", "COMPLETE_WITH_FAILURES"} for row in provenance),
            "non-terminal source run")
    require(all(
        row["arm"] == "safety_sentry"
        for row in provenance if row["status"] == "COMPLETE_WITH_FAILURES"
    ), "only superseded Safety Sentry shards may retain infrastructure failures")
    stale_sources = [row for row in provenance if row["source_manifest_status"] == "RUNNING"]
    require(
        len(stale_sources) == 1
        and stale_sources[0]["run_id"] == "r5v10_git_saf_b3"
        and stale_sources[0]["status"] == "COMPLETE_WITH_FAILURES"
        and stale_sources[0]["recovery_role"] == "superseded_interrupted_base",
        "unexpected stale top-level source status",
    )
    require(
        sum(row["recovery_role"] == "replacement" for row in provenance) == 3,
        "expected exactly three recovery shards",
    )
    require(len({row["measurement_code_fingerprint"] for row in provenance}) == 1,
            "measurement source fingerprints are not identical")

    checks = []
    for label, arm in METHODS.items():
        arm_sequences = [row for row in sequences if row["arm"] == arm]
        arm_tasks = [row for row in tasks if row["arm"] == arm]
        summary = next(row for row in paper if row["method"] == label)
        require(len(arm_sequences) == 32, f"{label}: expected 32 sequences")
        require(len(arm_tasks) == 128, f"{label}: expected 128 tasks")
        require(len({row["sequence_id"] for row in arm_sequences}) == 32,
                f"{label}: duplicate sequence rows")
        require(len({(row["sequence_id"], row["task_id"]) for row in arm_tasks}) == 128,
                f"{label}: duplicate task rows")
        task_rate = sum(bool(row["task_success"]) for row in arm_tasks) / 128
        future = [row for row in arm_tasks if int(row["task_index"]) > 1]
        future_rate = sum(bool(row["task_success"]) for row in future) / 96
        no_violation = sum(not bool(row["safety_violation"]) for row in future) / 96
        require(close(summary["task_success_rate"], task_rate), f"{label}: task rate mismatch")
        require(close(summary["future_task_success_rate"], future_rate),
                f"{label}: future success mismatch")
        require(close(summary["future_no_violation_rate"], no_violation),
                f"{label}: no-violation mismatch")
        checks.append({
            "method": label,
            "sequences": len(arm_sequences),
            "tasks": len(arm_tasks),
            "future_tasks": len(future),
            "future_task_success_rate": future_rate,
            "future_no_violation_rate": no_violation,
        })

    expected_baselines = set(METHODS.values()) - {METHODS["StateWise"]}
    require({row["baseline"] for row in paired_json} == expected_baselines,
            "paired JSON baseline set mismatch")
    require(len(paired_json) == len(paired_sheet) == 3, "paired result count mismatch")
    json_by_baseline = {row["baseline"]: row for row in paired_json}
    for row in paired_sheet:
        other = json_by_baseline[row["baseline"]]
        require(int(row["sequences"]) == int(other["sequences"]) == 32,
                f"{row['baseline']}: paired N mismatch")
        for key in ("paired_risk_difference", "mcnemar_p", "holm_adjusted_p"):
            require(close(row[key], other[key]), f"{row['baseline']}: {key} mismatch")

    run_counts = Counter(row["arm"] for row in provenance)
    expected_run_counts = {
        "statewise_v2_contract_corrected_typedfix": 4,
        "memlineage_trust": 4,
        "safety_sentry": 7,
        "ask_or_assume": 4,
    }
    require(all(run_counts[arm] == count for arm, count in expected_run_counts.items()),
            "paper arm provenance shard count mismatch")
    report = {
        "status": "VALID",
        "validated_at": datetime.now(timezone.utc).isoformat(),
        "workbook": str(WORKBOOK),
        "paired_json": str(PAIRED),
        "dataset_sha256": DATASET_SHA,
        "measurement_code_fingerprint": provenance[0]["measurement_code_fingerprint"],
        "paper_methods": checks,
        "paired_comparisons": paired_json,
    }
    REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(REPORT.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
