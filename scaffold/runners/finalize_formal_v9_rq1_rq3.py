from __future__ import annotations

from collections import defaultdict
from math import comb
from pathlib import Path
from random import Random
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'paper_experiments/results/rq2/rq2_results.xlsx'
OUTPUTS = {
    "RQ1": ROOT / 'paper_experiments/results/rq1/rq1_results.xlsx',
    "RQ3": ROOT / 'paper_experiments/results/rq3/rq3_results.xlsx',
}
ARMS = {
    "RQ1": [
        "correct",
        "corrupted",
        "no_memory",
        "statewise_v2_contract_corrected_typedfix",
    ],
    "RQ3": [
        "correct",
        "corrupted",
        "no_memory",
        "memlineage_trust",
        "safety_sentry",
        "ask_or_assume",
        "statewise_v2_contract_corrected_typedfix",
    ],
}
LABELS = {
    "correct": "Correct memory",
    "corrupted": "Corrupted memory",
    "no_memory": "No memory",
    "memlineage_trust": "MemLineage",
    "safety_sentry": "Safety Sentry",
    "ask_or_assume": "Ask-or-Assume",
    "statewise_v2_contract_corrected_typedfix": "StateWise",
}


def sheet_records(workbook, sheet_name: str) -> tuple[list[Any], list[dict[str, Any]]]:
    values = list(workbook[sheet_name].values)
    headers = list(values[0])
    return headers, [dict(zip(headers, row)) for row in values[1:]]


def exact_mcnemar(left_only: int, right_only: int) -> float:
    discordant = left_only + right_only
    if discordant == 0:
        return 1.0
    tail = sum(comb(discordant, k) for k in range(min(left_only, right_only) + 1))
    return min(1.0, 2.0 * tail / (2**discordant))


def paired_bootstrap_ci(left: list[bool], right: list[bool]) -> tuple[float, float]:
    random = Random(20260815)
    n = len(left)
    effects = []
    for _ in range(10_000):
        indices = [random.randrange(n) for _ in range(n)]
        effects.append(sum(left[index] - right[index] for index in indices) / n)
    effects.sort()
    return effects[249], effects[9749]


def holm_adjust(p_values: list[float]) -> list[float]:
    indexed = sorted(enumerate(p_values), key=lambda item: item[1])
    adjusted = [1.0] * len(p_values)
    running = 0.0
    total = len(p_values)
    for rank, (index, p_value) in enumerate(indexed):
        running = max(running, min(1.0, p_value * (total - rank)))
        adjusted[index] = running
    return adjusted


def style(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="35546D")
        for column in sheet.columns:
            width = min(52, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width


def write_records(sheet, headers: list[Any], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])


def build(rq: str, source) -> None:
    selected = set(ARMS[rq])
    summary_headers, all_summary = sheet_records(source, "Arm_Summary")
    episode_headers, all_episodes = sheet_records(source, "Episode_Data")
    failure_headers, all_failures = sheet_records(source, "Failure_Audit")
    run_headers, run_rows = sheet_records(source, "Run_Status")

    summaries = [row for row in all_summary if row["arm"] in selected]
    episodes = [row for row in all_episodes if row["arm"] in selected]
    failures = [row for row in all_failures if row["arm"] in selected]

    workbook = Workbook()
    paper = workbook.active
    paper.title = "Paper_Table"
    paper_headers = [
        "method", "internal_arm", "n", "objective_correct_n",
        "objective_correctness", "should_act_n", "action_completed_n",
        "action_completion", "should_stop_n", "correct_safe_stop_n",
        "safe_stop", "safety_violation_n", "safety_violation",
        "unnecessary_mutation_rate",
    ]
    paper.append(paper_headers)
    by_arm = {row["arm"]: row for row in summaries}
    for arm in ARMS[rq]:
        row = by_arm[arm]
        paper.append([
            LABELS[arm], arm, row["episodes_succeeded"],
            row["behaviorally_correct_n"], row["behavioral_correctness_rate"],
            row["should_act_n"], row["act_correct_n"], row["action_completion_rate"],
            row["should_stop_n"], row["stop_correct_n"], row["safe_stop_rate"],
            row["safety_violation_n"], row["safety_violation_rate"],
            row["mean_unnecessary_mutation_rate"],
        ])

    write_records(workbook.create_sheet("Arm_Summary"), summary_headers, summaries)
    write_records(workbook.create_sheet("Episode_Data"), episode_headers, episodes)

    domain = workbook.create_sheet("Domain_Summary")
    domain_headers = [
        "method", "internal_arm", "domain", "n", "objective_correct_n",
        "objective_correctness", "safety_violation_n", "safety_violation",
    ]
    domain.append(domain_headers)
    buckets: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in episodes:
        if row["episode_status"] == "SUCCEEDED":
            buckets[(row["arm"], row["domain"])].append(row)
    for arm in ARMS[rq]:
        for (bucket_arm, domain_name), rows in sorted(buckets.items()):
            if bucket_arm != arm:
                continue
            n = len(rows)
            correct = sum(bool(row["behaviorally_correct"]) for row in rows)
            violations = sum(bool(row["safety_violation"]) for row in rows)
            domain.append([
                LABELS[arm], arm, domain_name, n, correct, correct / n,
                violations, violations / n,
            ])

    paired = workbook.create_sheet("Paired_Contrasts")
    paired_headers = [
        "left", "right", "metric", "paired_cases", "left_rate", "right_rate",
        "risk_difference", "bootstrap_ci95_low", "bootstrap_ci95_high",
        "left_only", "right_only", "exact_mcnemar_p", "holm_adjusted_p",
    ]
    paired.append(paired_headers)
    comparisons = (
        [("correct", "corrupted"), ("no_memory", "corrupted")]
        if rq == "RQ1"
        else [("statewise_v2_contract_corrected_typedfix", arm) for arm in ARMS[rq][:-1]]
    )
    outcome = {
        (row["arm"], row["case_id"]): bool(row["behaviorally_correct"])
        for row in episodes if row["episode_status"] == "SUCCEEDED"
    }
    paired_rows = []
    for left, right in comparisons:
        cases = sorted(
            {case for arm, case in outcome if arm == left}
            & {case for arm, case in outcome if arm == right}
        )
        left_values = [outcome[(left, case)] for case in cases]
        right_values = [outcome[(right, case)] for case in cases]
        left_only = sum(l and not r for l, r in zip(left_values, right_values))
        right_only = sum(r and not l for l, r in zip(left_values, right_values))
        ci_low, ci_high = paired_bootstrap_ci(left_values, right_values)
        paired_rows.append({
            "left": LABELS[left], "right": LABELS[right],
            "metric": "objective_correctness", "paired_cases": len(cases),
            "left_rate": sum(left_values) / len(cases),
            "right_rate": sum(right_values) / len(cases),
            "risk_difference": (sum(left_values) - sum(right_values)) / len(cases),
            "bootstrap_ci95_low": ci_low, "bootstrap_ci95_high": ci_high,
            "left_only": left_only, "right_only": right_only,
            "exact_mcnemar_p": exact_mcnemar(left_only, right_only),
        })
    adjusted = holm_adjust([row["exact_mcnemar_p"] for row in paired_rows])
    for row, p_holm in zip(paired_rows, adjusted):
        row["holm_adjusted_p"] = p_holm
        paired.append([row.get(header) for header in paired_headers])

    plot = workbook.create_sheet("Plot_Data")
    plot.append(["rq", "method", "metric", "value", "numerator", "denominator"])
    for arm in ARMS[rq]:
        row = by_arm[arm]
        for metric, value, numerator, denominator in [
            ("objective_correctness", row["behavioral_correctness_rate"], row["behaviorally_correct_n"], row["episodes_succeeded"]),
            ("action_completion", row["action_completion_rate"], row["act_correct_n"], row["should_act_n"]),
            ("safe_stop", row["safe_stop_rate"], row["stop_correct_n"], row["should_stop_n"]),
            ("safety_violation", row["safety_violation_rate"], row["safety_violation_n"], row["episodes_succeeded"]),
        ]:
            plot.append([rq, LABELS[arm], metric, value, numerator, denominator])

    write_records(workbook.create_sheet("Failures"), failure_headers, failures)
    write_records(workbook.create_sheet("Run_Provenance"), run_headers, run_rows)
    notes = workbook.create_sheet("Notes")
    notes.append(["field", "value"])
    notes.append(["source", str(SOURCE.relative_to(ROOT))])
    notes.append(["measurement_contract", "Corrected formal-v9 objective evaluator; arbitrary safe refusal is not task success."])
    notes.append(["dataset_snapshot", "w1_contract_corrected_v9"])
    notes.append(["dataset_sha256", "25f08c7ec7f5cee5855dcddbfc9e538bab79a509aa8580dbd7501ef396c2af34"])
    notes.append(["oracle", "Excluded from paper-facing workbooks."])
    notes.append(["status", "Extended validation on a fixed evaluation suite; not untouched held-out evidence."])
    style(workbook)
    OUTPUTS[rq].parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUTS[rq])
    print(OUTPUTS[rq].resolve())


def main() -> int:
    source = load_workbook(SOURCE, read_only=True, data_only=True)
    for rq in OUTPUTS:
        build(rq, source)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
