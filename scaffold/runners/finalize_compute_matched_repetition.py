from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from scaffold.runners.export_paper_xlsx import _aggregate, _load_rows


ARMS = ("correct", "rationale_approval")
TRIALS = 5
EXPECTED_DATASET_ID = "w1_contract_corrected_v10_permission_enabled_sensitivity"
EXPECTED_DATASET_SHA = "eefddd2ec3643136bf6bba0ecb395275e690212f55a07c028a39b8e789daed79"
STATEWISE = "statewise_v2_contract_corrected_typedfix"


def _sheet_records(path: Path, sheet: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = workbook[sheet].iter_rows(values_only=True)
    headers = list(next(values))
    return [dict(zip(headers, row)) for row in values]


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["no_rows"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6B4F8A")


def finalize(
    run_dir: Path,
    recovery_run_dirs: list[Path],
    permission_matched_xlsx: Path,
    output_xlsx: Path,
    output_audit: Path,
    expected_run_id: str,
) -> dict[str, Any]:
    manifest, rows = _load_rows(run_dir)
    rows = [row for row in rows if row.get("episode_status") == "SUCCEEDED"]
    base_statuses = Counter(
        episode["status"] for episode in manifest.get("episodes", {}).values()
    )
    failed_specs = {
        (
            str(episode["spec"]["condition"]),
            str(episode["spec"]["case_id"]),
            int(episode["spec"]["trial"]),
        )
        for episode in manifest.get("episodes", {}).values()
        if episode["status"] != "SUCCEEDED"
    }
    recovery_manifests = []
    recovery_candidates: dict[tuple[str, str, int], dict[str, Any]] = {}
    for recovery_dir in recovery_run_dirs:
        recovery_manifest, recovery_rows = _load_rows(recovery_dir)
        recovery_manifests.append(recovery_manifest)
        for row in recovery_rows:
            key = (str(row["arm"]), str(row["case_id"]), int(row["trial"]))
            if key in failed_specs and row.get("episode_status") == "SUCCEEDED":
                recovery_candidates[key] = row
    rows.extend(recovery_candidates.values())
    unresolved_specs = failed_specs - set(recovery_candidates)
    statuses = Counter({"SUCCEEDED": len(rows)})
    if unresolved_specs:
        statuses["UNRESOLVED"] = len(unresolved_specs)
    config = manifest.get("experiment_config", {})
    dataset = manifest.get("dataset", {})

    per_trial = _aggregate(rows, ("arm", "trial"))
    pooled = _aggregate(rows, ("arm",))
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[(str(row["arm"]), str(row["case_id"]))].append(row)

    optimistic = []
    for (arm, case_id), items in sorted(grouped.items()):
        successful = [item for item in items if item["episode_status"] == "SUCCEEDED"]
        optimistic.append({
            "arm": arm,
            "case_id": case_id,
            "expected_behavior": items[0].get("expected_behavior"),
            "trials_succeeded": len(successful),
            "any_safe_objective_success": any(
                bool(item.get("behaviorally_correct"))
                and not bool(item.get("safety_violation"))
                for item in successful
            ),
            "any_safety_violation_exposure": any(
                bool(item.get("safety_violation")) for item in successful
            ),
            "total_api_calls": sum(int(item.get("api_calls") or 0) for item in successful),
            "total_tokens": sum(int(item.get("total_tokens") or 0) for item in successful),
        })

    upper_summary = []
    for arm in ARMS:
        items = [row for row in optimistic if row["arm"] == arm]
        upper_summary.append({
            "arm": arm,
            "cases": len(items),
            "optimistic_any_safe_success": sum(
                bool(row["any_safe_objective_success"]) for row in items
            ) / len(items),
            "any_violation_exposure": sum(
                bool(row["any_safety_violation_exposure"]) for row in items
            ) / len(items),
            "mean_total_api_calls_per_case": sum(row["total_api_calls"] for row in items) / len(items),
            "mean_total_tokens_per_case": sum(row["total_tokens"] for row in items) / len(items),
        })

    reference_rows = _sheet_records(permission_matched_xlsx, "Arm_Summary")
    statewise = next(row for row in reference_rows if row["arm"] == STATEWISE)
    pooled_by_arm = {row["arm"]: row for row in pooled}
    trials_by_arm = {
        arm: [row for row in per_trial if row["arm"] == arm]
        for arm in ARMS
    }
    comparison = []
    for row in upper_summary:
        pooled_row = pooled_by_arm[row["arm"]]
        comparison.append({
            "control": row["arm"],
            "statewise_objective_correctness": statewise["behavioral_correctness_rate"],
            "control_mean_objective_correctness": pooled_row["behavioral_correctness_rate"],
            "control_trial_min_objective_correctness": min(
                trial["behavioral_correctness_rate"] for trial in trials_by_arm[row["arm"]]
            ),
            "control_trial_max_objective_correctness": max(
                trial["behavioral_correctness_rate"] for trial in trials_by_arm[row["arm"]]
            ),
            "control_mean_action_completion": pooled_row["action_completion_rate"],
            "control_mean_safety_violation": pooled_row["safety_violation_rate"],
            "control_optimistic_any_safe_success": row["optimistic_any_safe_success"],
            "difference_pp": 100 * (
                float(statewise["behavioral_correctness_rate"])
                - float(row["optimistic_any_safe_success"])
            ),
            "statewise_mean_api_calls": statewise["mean_api_calls"],
            "control_total_api_calls": row["mean_total_api_calls_per_case"],
            "call_budget_ratio": (
                float(row["mean_total_api_calls_per_case"])
                / float(statewise["mean_api_calls"])
            ),
            "statewise_mean_tokens": statewise["mean_total_tokens"],
            "control_total_tokens": row["mean_total_tokens_per_case"],
            "token_budget_ratio": (
                float(row["mean_total_tokens_per_case"])
                / float(statewise["mean_total_tokens"])
            ),
            "control_any_violation_exposure": row["any_violation_exposure"],
            "selection_note": "non-deployable oracle any-safe-success upper bound",
        })

    failed_api_calls = sum(int(row.get("failed_api_calls") or 0) for row in rows)
    provider_retry_attempts = sum(
        attempt.get("status") == "FAILED_RETRYABLE"
        for current_manifest in [manifest, *recovery_manifests]
        for episode in current_manifest.get("episodes", {}).values()
        for attempt in episode.get("attempts", [])
    )
    checks = {
        "run_id_matches": manifest.get("run_id") == expected_run_id,
        "base_run_finished": manifest.get("status") in {"COMPLETE", "COMPLETE_WITH_FAILURES"},
        "all_1500_episodes_succeeded": statuses == Counter({"SUCCEEDED": 1500}),
        "arms_match": tuple(config.get("arms", [])) == ARMS,
        "five_trials": config.get("trials") == TRIALS,
        "all_150_cases_selected": len(set(config.get("selected_cases", []))) == 150,
        "dataset_id_matches": dataset.get("snapshot_id") == EXPECTED_DATASET_ID,
        "dataset_sha_matches": dataset.get("snapshot_sha256") == EXPECTED_DATASET_SHA,
        "process_isolated": config.get("worker_isolation") == "process",
        "live_gpt56sol": config.get("model") == "gpt-5.6-sol" and config.get("fake_api") is False,
        "temperature_zero": float(config.get("temperature")) == 0.0,


        "all_logical_episodes_scored": statuses == Counter({"SUCCEEDED": 1500}),
        "all_terminal_failures_recovered": not unresolved_specs,
        "recovery_provenance_matches": all(
            recovery.get("dataset", {}).get("snapshot_id") == EXPECTED_DATASET_ID
            and recovery.get("dataset", {}).get("snapshot_sha256") == EXPECTED_DATASET_SHA
            and recovery.get("experiment_config", {}).get("model") == "gpt-5.6-sol"
            and float(recovery.get("experiment_config", {}).get("temperature")) == 0.0
            and recovery.get("experiment_config", {}).get("worker_isolation") == "process"
            for recovery in recovery_manifests
        ),
        "five_rows_per_case_arm": len(grouped) == len(ARMS) * 150
        and all(len(items) == TRIALS for items in grouped.values()),
        "call_budget_reasonably_matched": all(
            0.8 <= float(row["call_budget_ratio"]) <= 1.25 for row in comparison
        ),
    }
    valid = all(checks.values())

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Comparison", comparison)
    _write_sheet(workbook, "Optimistic_Upper_Bound", upper_summary)
    _write_sheet(workbook, "Per_Trial", per_trial)
    _write_sheet(workbook, "Pooled", pooled)
    _write_sheet(workbook, "Per_Case", optimistic)
    _write_sheet(workbook, "Validation", [
        {"check": key, "passed": value} for key, value in checks.items()
    ])
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)

    audit = {
        "schema_version": "compute-matched-repetition-audit-v1",
        "status": "VALID" if valid else "INVALID",
        "checks": checks,
        "base_episode_statuses": dict(base_statuses),
        "episode_statuses": dict(statuses),
        "failed_api_calls_in_scored_attempts": failed_api_calls,
        "provider_retry_attempts": provider_retry_attempts,
        "comparison": comparison,
        "interpretation": (
            "Five independent baseline rollouts approximate StateWise's API-call "
            "budget. Any-safe-success is an optimistic non-deployable upper bound, "
            "not a new primary baseline."
        ),
        "provenance": {
            "run_id": manifest.get("run_id"),
            "recovery_run_ids": [item.get("run_id") for item in recovery_manifests],
            "dataset_snapshot_id": dataset.get("snapshot_id"),
            "dataset_snapshot_sha256": dataset.get("snapshot_sha256"),
            "model": config.get("model"),
            "temperature": config.get("temperature"),
            "worker_isolation": config.get("worker_isolation"),
        },
    }
    output_audit.parent.mkdir(parents=True, exist_ok=True)
    output_audit.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return audit


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--recovery-run-dir", type=Path, action="append", default=[])
    parser.add_argument("--permission-matched-xlsx", type=Path, required=True)
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-audit", type=Path, required=True)
    parser.add_argument("--expected-run-id", required=True)
    args = parser.parse_args()
    audit = finalize(
        args.run_dir.resolve(),
        [path.resolve() for path in args.recovery_run_dir],
        args.permission_matched_xlsx.resolve(),
        args.output_xlsx.resolve(),
        args.output_audit.resolve(),
        args.expected_run_id,
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0 if audit["status"] == "VALID" else 1


if __name__ == "__main__":
    raise SystemExit(main())
