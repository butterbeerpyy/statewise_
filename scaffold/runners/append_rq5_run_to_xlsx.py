from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
RUN = ROOT / "runs" / "rq5_longitudinal" / "rq5_memlineage_gpt56sol_recheck_20260810"
BOOK = ROOT / 'paper_experiments' / 'results' / "rq5" / "rq5_results.xlsx"
RUN_ID = "rq5_memlineage_gpt56sol_recheck_20260810"


def rows_from_summary(summary: dict) -> list[dict]:
    rows = []
    for item in summary.get("sequence_arm_results", []):
        seq = str(item["sequence_id"])
        parts = seq.split("_")
        domain = parts[1] if len(parts) >= 3 else None
        index = int(parts[2]) if len(parts) >= 3 and parts[2].isdigit() else None
        rows.append({
            "run_id": RUN_ID,
            "model": "gpt-5.6-sol",
            "sequence_id": seq,
            "domain": domain,
            "sequence_index": index,
            "arm": item.get("arm"),
            "trial": item.get("trial"),
            "tasks": item.get("tasks"),
            "task_success_count": item.get("task_success_count"),
            "task_success_rate": item.get("task_success_rate"),
            "future_task_success_count": item.get("future_task_success_count"),
            "future_task_success_rate": item.get("future_task_success_rate"),
            "future_task_safety_rate": item.get("future_task_safety_rate"),
            "state_handoff_integrity_rate": item.get("state_handoff_integrity_rate"),
            "state_reuse_rate": item.get("state_reuse_rate"),
            "repeated_intervention_rate": item.get("repeated_intervention_rate"),
            "cumulative_questions": item.get("cumulative_questions"),
            "cumulative_permission_interactions": item.get("cumulative_permission_interactions"),
            "cumulative_safety_violations": item.get("cumulative_safety_violations"),
            "t1_repair_success": item.get("t1_repaired_contract_satisfied"),
            "primary_in_scope": True,
            "exclusion_reason": None,
        })
    return rows


def append_rows(ws, rows: list[dict], key: str = "run_id") -> None:
    headers = [c.value for c in ws[1]]
    existing = [r for r in ws.iter_rows(min_row=2, values_only=True)
                if not (r[headers.index(key)] == RUN_ID)]
    ws.delete_rows(2, ws.max_row)
    for old in existing:
        ws.append(list(old))
    for row in rows:
        ws.append([row.get(h) for h in headers])


def main() -> None:
    summary = json.loads((RUN / "longitudinal_summary.json").read_text(encoding="utf-8"))
    rows = rows_from_summary(summary)
    wb = load_workbook(BOOK)
    append_rows(wb["Primary_Sequence_Data"], rows)
    ws = wb["Primary_Summary"]
    headers = [c.value for c in ws[1]]
    ws.delete_rows(2, ws.max_row)
    for old in [("gpt-5.6-sol", "statewise_v2_retrieval", 16, 64, 1, 1, 1, 1, 1, 0, 1)]:
        ws.append(list(old))
    table = summary["table"][0]
    summary_row = [table.get(h) if h not in {"model", "tasks"} else ("gpt-5.6-sol" if h == "model" else 64) for h in headers]
    ws.append(summary_row)
    ws = wb["Run_Status"]
    headers = [c.value for c in ws[1]]
    run_col = headers.index("run_id")
    old_status = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[run_col] != RUN_ID]
    ws.delete_rows(2, ws.max_row)
    for old in old_status:
        ws.append(list(old))
    ws.append([RUN_ID, "gpt-5.6-sol", "COMPLETE", 16, 16, 0, 0, True, str(RUN / "manifest.json")])
    wb.save(BOOK)


if __name__ == "__main__":
    main()
