from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
V2_ROOT = ROOT.parent / 'statewise'
DEFAULT_DEST = ROOT / 'paper_experiments' / 'results'
DEFAULT_RQ14 = ROOT / 'paper_experiments' / "analysis" / "rq1_rq4_crossmodel20_paper_data_FINAL.xlsx"
SCOPE_PATH = V2_ROOT / "scaffold" / "experiments" / "rq5_analysis_scope_v1.json"
RQ5_ROOT = ROOT / "runs" / "rq5_longitudinal"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_sheet(ws, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["no rows"])
        return
    headers = list(rows[0])
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        ws.append([row.get(h) for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, len(headers) + 1):
        vals = [str(ws.cell(r, col).value or "") for r in range(1, min(100, ws.max_row) + 1)]
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(12, max(map(len, vals), default=10) + 2))


def values_from_sheet(src, sheet_name: str) -> list[dict[str, Any]]:
    ws = src[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def export_single_sheet(src_path: Path, sheet_name: str, out_path: Path) -> None:
    src = load_workbook(src_path, read_only=True, data_only=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    write_sheet(ws, values_from_sheet(src, sheet_name))
    notes = wb.create_sheet("Scope_Notes")
    write_sheet(notes, [{"item": "source_workbook", "value": str(src_path)}, {"item": "source_sheet", "value": sheet_name}])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def sequence_domain_index(sequence_id: str) -> tuple[str | None, int | None]:
    match = re.match(r"seq_([^_]+)_([0-9]+)$", sequence_id)
    if not match:
        return None, None
    return match.group(1), int(match.group(2))


def rq5_rows() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows = []
    source_specs = [
        ("rq5_v2_full40_gpt56sol", "gpt-5.6-sol"),
        ("rq5_v2_full40_gpt56sol_part2", "gpt-5.6-sol"),
    ]
    included = {"environment", "git", "database", "deployment"}
    for run_id, model in source_specs:
        path = RQ5_ROOT / run_id / "longitudinal_summary.json"
        if not path.exists():
            continue
        summary = read_json(path)
        for item in summary.get("sequence_arm_results", []):
            seq = item.get("sequence_id")
            domain, index = sequence_domain_index(str(seq))
            in_scope = bool(domain in included and index is not None and index <= 4)
            reason = None
            if domain == "filesystem" and index is not None and index <= 4:
                reason = "method_limitation_pre_registered_exclusion"
            elif domain not in included:
                reason = "domain_not_in_primary_scope"
            elif index is None or index > 4:
                reason = "outside_four_sequences_per_domain_primary_quota"
            rows.append({
                "run_id": run_id,
                "model": model,
                "sequence_id": seq,
                "domain": domain,
                "sequence_index": index,
                "arm": item.get("arm"),
                "trial": item.get("trial"),
                "tasks": item.get("tasks"),
                "task_success_count": item.get("task_success_count"),
                "task_success_rate": item.get("task_success_rate"),
                "future_task_success_count": item.get("future_task_success_count"),
                "future_task_success_rate": item.get("future_task_success_rate"),
                "future_task_safety_rate": item.get("future_task_safety_rate"),
                "state_handoff_integrity_rate": item.get("state_handoff_integrity_rate"),
                "state_reuse_rate": item.get("state_reuse_rate"),
                "repeated_intervention_rate": item.get("repeated_intervention_rate"),
                "cumulative_questions": item.get("cumulative_questions"),
                "cumulative_permission_interactions": item.get("cumulative_permission_interactions"),
                "cumulative_safety_violations": item.get("cumulative_safety_violations"),
                "t1_repair_success": item.get("t1_repaired_contract_satisfied"),
                "primary_in_scope": in_scope,
                "exclusion_reason": reason,
            })
    return [r for r in rows if r["primary_in_scope"]], rows


def rq5_primary_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups = {}
    for row in rows:
        key = (row["model"], row["arm"])
        groups.setdefault(key, []).append(row)
    out = []
    for (model, arm), items in sorted(groups.items()):
        n = len(items)
        future_den = sum(3 for _ in items)
        out.append({
            "model": model,
            "arm": arm,
            "sequences": n,
            "tasks": sum(int(x.get("tasks") or 0) for x in items),
            "task_success_rate": sum(float(x.get("task_success_rate") or 0) for x in items) / n if n else None,
            "future_task_success_rate": sum(int(round(float(x.get("future_task_success_rate") or 0) * 3)) for x in items) / future_den if future_den else None,
            "future_task_safety_rate": sum(float(x.get("future_task_safety_rate") or 0) for x in items) / n if n else None,
            "state_handoff_integrity_rate": sum(float(x.get("state_handoff_integrity_rate") or 0) for x in items) / n if n else None,
            "state_reuse_rate": sum(float(x.get("state_reuse_rate") or 0) for x in items) / n if n else None,
            "repeated_intervention_rate": sum(float(x.get("repeated_intervention_rate") or 0) for x in items) / n if n else None,
            "t1_repair_success_rate": sum(bool(x.get("t1_repair_success")) for x in items) / n if n else None,
        })
    return out


def rq5_run_status() -> list[dict[str, Any]]:
    rows = []
    for run_id, model in [
        ("rq5_v2_full40_gpt56sol", "gpt-5.6-sol"),
        ("rq5_v2_full40_gpt56sol_part2", "gpt-5.6-sol"),
        ("rq5v2_method_deepseek", "deepseek-v4-pro"),
        ("rq5v2_method_qwen", "qwen3.7-max"),
    ]:
        manifest_path = RQ5_ROOT / run_id / "manifest.json"
        summary_path = RQ5_ROOT / run_id / "longitudinal_summary.json"
        if not manifest_path.exists():
            continue
        manifest = read_json(manifest_path)
        eps = manifest.get("episodes") or {}
        counts = Counter(x.get("status") for x in eps.values())
        rows.append({
            "run_id": run_id,
            "model": model,
            "status": manifest.get("status"),
            "episodes": len(eps),
            "succeeded": counts.get("SUCCEEDED", 0),
            "running": counts.get("RUNNING", 0),
            "pending": counts.get("PENDING", 0),
            "summary_available": summary_path.exists(),
            "source_manifest": str(manifest_path),
        })
    rows.append({
        "run_id": "rq5_claude_reported_no_local_manifest",
        "model": "claude-fable-5",
        "status": "REPORTED_EXTERNALLY_MANIFEST_NOT_FOUND",
        "episodes": None,
        "succeeded": None,
        "running": None,
        "pending": None,
        "summary_available": False,
        "source_manifest": None,
    })
    return rows


def build_audit_workbook(index_path: Path, out_path: Path) -> None:
    index = read_json(index_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Run_Metadata"
    metadata = [{"field": key, "value": json.dumps(index.get(key), ensure_ascii=False) if isinstance(index.get(key), (dict, list)) else index.get(key)} for key in [
        "source_run", "run_status", "model", "api_profile", "fake_api", "arms", "trials",
        "attribution_policy", "total_episodes", "scored_episodes", "not_yet_run_episodes", "exported",
    ]]
    metadata.extend([
        {"field": "publication_class", "value": "secondary_audit_only"},
        {"field": "exclusion_reason", "value": "historical run has known measurement/runtime contamination or incomplete coverage; not a headline result"},
    ])
    write_sheet(ws, metadata)
    bundles = []
    for item in index.get("bundles", []):
        row = {k: v for k, v in item.items() if not isinstance(v, (dict, list))}
        for key in ("difficulty",):
            if key in item:
                row[key] = json.dumps(item[key], ensure_ascii=False)
        bundles.append(row)
    ws = wb.create_sheet("Bundle_Data")
    write_sheet(ws, bundles)
    ws = wb.create_sheet("Failure_Classes")
    write_sheet(ws, [{"failure_class": k, "count": v} for k, v in sorted((index.get("failure_classes") or {}).items())])
    ws = wb.create_sheet("Symptoms")
    write_sheet(ws, [{"symptom": k, "count": v} for k, v in sorted((index.get("symptom_counts") or {}).items())])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def build_recovery_workbook(analysis_paths: list[Path], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Run_Summary"
    summaries = []
    for path in analysis_paths:
        if not path.exists():
            continue
        report = read_json(path)
        summaries.append({k: report.get(k) for k in [
            "run_id", "run_status", "model", "arm", "worker_isolation", "episodes",
            "execution_task_success", "acceptable_safe_terminal", "objective_behaviorally_correct",
            "safety_violations", "repair_events", "repair_correct_events", "primary_candidate_seen",
            "primary_critical_seen", "execution_isolation_valid",
        ]})
    write_sheet(ws, summaries)
    ws = wb.create_sheet("Episode_Diagnosis")
    all_rows = []
    for path in analysis_paths:
        if not path.exists():
            continue
        report = read_json(path)
        for row in report.get("rows", []):
            all_rows.append({"run_id": report.get("run_id"), "model": report.get("model"), **{k: v for k, v in row.items() if not isinstance(v, (dict, list))}})
    write_sheet(ws, all_rows)
    ws = wb.create_sheet("Scope_Notes")
    write_sheet(ws, [{"item": "publication_class", "value": "secondary corrected-v2 recovery; not the formal held-out 300"}, {"item": "reason", "value": "53-case recovery is useful engineering evidence but not a replacement for the pre-registered headline matrix"}])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def build_formal_slot_inventory(out_path: Path) -> None:
    slots_path = ROOT / "dataset_design" / "artifacts" / "paper_benchmark_v1" / "paper_heldout_300_slots_v1.json"
    registry_path = ROOT / "dataset_design" / "artifacts" / "paper_heldout_materialization_v1" / "materialization_registry.json"
    slots = read_json(slots_path).get("slots", [])
    registry = {item.get("slot_id"): item for item in read_json(registry_path).get("cases", [])}
    rows = []
    for slot in slots:
        reg = registry.get(slot.get("slot_id"), {})
        rows.append({
            "slot_id": slot.get("slot_id"),
            "domain": slot.get("operational_domain"),
            "selection_role": reg.get("selection_role"),
            "materialization_status": reg.get("materialization_status"),
            "corruption_operator": slot.get("corruption_operator"),
            "causal_depth": slot.get("causal_depth"),
            "observability": slot.get("observability"),
            "impact": slot.get("impact"),
            "state_cardinality": slot.get("state_cardinality"),
            "task_contract": slot.get("task_contract"),
            "case_id": reg.get("case_id"),
            "schema_valid": reg.get("schema_valid"),
            "deterministic_closure_valid": reg.get("deterministic_closure_valid"),
            "annotation_a_status": reg.get("annotation_a_status"),
            "annotation_b_status": reg.get("annotation_b_status"),
            "adjudication_status": reg.get("adjudication_status"),
        })
    wb = Workbook()
    ws = wb.active
    ws.title = "Slot_Inventory"
    write_sheet(ws, rows)
    ws = wb.create_sheet("Counts")
    from collections import Counter
    write_sheet(ws, [{"field": "selection_role", "value": k, "count": v} for k, v in Counter(r["selection_role"] for r in rows).items()] + [{"field": "materialization_status", "value": k, "count": v} for k, v in Counter(r["materialization_status"] for r in rows).items()])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def build(dest: Path, rq14_path: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    for name in ("rq1", "rq2", "rq3", "rq4", "rq5"):
        (dest / name).mkdir(exist_ok=True)
    shutil.copy2(rq14_path, dest / "rq1_rq4_crossmodel20_paper_data_FINAL.xlsx")
    src = load_workbook(rq14_path, read_only=True, data_only=True)
    for i in range(1, 5):
        export_single_sheet(rq14_path, f"RQ{i}_Validation", dest / f"rq{i}" / f"RQ{i}_validation.xlsx")

    scope = read_json(SCOPE_PATH)
    primary, all_rows = rq5_rows()
    wb5 = Workbook()
    ws = wb5.active
    ws.title = "Scope"
    write_sheet(ws, [{"field": k, "value": json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in scope.items()])
    ws = wb5.create_sheet("Primary_Sequence_Data")
    write_sheet(ws, primary)
    ws = wb5.create_sheet("Primary_Summary")
    write_sheet(ws, rq5_primary_summary(primary))
    ws = wb5.create_sheet("Excluded_Sequences")
    write_sheet(ws, [r for r in all_rows if not r["primary_in_scope"]])
    ws = wb5.create_sheet("Run_Status")
    write_sheet(ws, rq5_run_status())
    ws = wb5.create_sheet("Notes")
    write_sheet(ws, [
        {"item": "primary", "value": "16 sequences: environment/git/database/deployment, 4 per domain"},
        {"item": "filesystem", "value": "4 filesystem sequences excluded by pre-registered method limitation; never treat as benchmark/runtime defect"},
        {"item": "limitations", "value": "seq_database_004 and seq_git_001 retained with ambiguity flags"},
        {"item": "annotation", "value": "Annotator A/B agreement 420/420; agreement does not establish benchmark validity"},
        {"item": "claude", "value": "User-reported Claude audit recorded, but raw local manifest was not found; no numeric Claude RQ5 result is fabricated"},
    ])
    wb5.save(dest / "rq5" / "RQ5_longitudinal_primary_and_audit.xlsx")

    secondary = dest / "secondary" / "audit"
    secondary.mkdir(parents=True, exist_ok=True)
    gpt55_index = ROOT / 'paper_experiments' / "for_debug" / "full224-statewise-v1-gpt55-20260801-r1" / "index.json"
    gpt56_index = ROOT / 'paper_experiments' / "for_debug" / "v1-full224-gpt56sol-validation" / "index.json"
    if gpt55_index.exists():
        build_audit_workbook(gpt55_index, secondary / "full224_statewise_v1_gpt55_audit.xlsx")
        shutil.copy2(gpt55_index, secondary / "full224_statewise_v1_gpt55_index.json")
    if gpt56_index.exists():
        build_audit_workbook(gpt56_index, secondary / "full224_statewise_v1_gpt56sol_partial_audit.xlsx")
        shutil.copy2(gpt56_index, secondary / "full224_statewise_v1_gpt56sol_partial_index.json")
    build_recovery_workbook([
        Path('runs/v2-typedfix-gpt56sol-recovery53-20260806-r1\\corrected_analysis.json'),
        Path('runs/v2-typedfix-deepseek-v4pro-recovery53-20260806-r1\\corrected_analysis.json'),
    ], dest / "secondary" / "corrected_v2_recovery53.xlsx")
    build_formal_slot_inventory(dest / "formal_300_slot_inventory.xlsx")


    master = load_workbook(rq14_path)
    for sheet in ("RQ5_Scope", "RQ5_Primary", "RQ5_Summary", "RQ5_Excluded", "RQ5_Run_Status", "RQ5_Notes"):
        if sheet in master.sheetnames:
            del master[sheet]
    for sheet, rows in [
        ("RQ5_Scope", [{"field": k, "value": json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v} for k, v in scope.items()]),
        ("RQ5_Primary", primary),
        ("RQ5_Summary", rq5_primary_summary(primary)),
        ("RQ5_Excluded", [r for r in all_rows if not r["primary_in_scope"]]),
        ("RQ5_Run_Status", rq5_run_status()),
    ]:
        ws = master.create_sheet(sheet)
        write_sheet(ws, rows)
    ws = master.create_sheet("RQ5_Notes")
    write_sheet(ws, [{"item": "scope", "value": "Primary four-domain, 16-sequence analysis"}, {"item": "filesystem", "value": "Pre-registered method limitation exclusion"}])
    master.save(dest / "RQ1_RQ5_RESULTS_MASTER.xlsx")

    (dest / "README.md").write_text(
        "# Paper experiment results\n\n"
        "This directory is organized by RQ1-RQ5. RQ1-RQ4 are the frozen 20-case cross-model validation matrix. "
        "RQ5 primary uses the pre-registered four-domain/16-sequence scope. Filesystem is retained only in the exclusion/audit sheet.\n\n"
        "The Claude RQ5 report supplied in conversation is recorded as an external report only; no local raw manifest was found, so no unverified numeric Claude RQ5 result is included.\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dest", type=Path, default=DEFAULT_DEST)
    parser.add_argument("--rq14", type=Path, default=DEFAULT_RQ14)
    args = parser.parse_args()
    build(args.dest, args.rq14)
    print(args.dest.resolve())


if __name__ == "__main__":
    main()
