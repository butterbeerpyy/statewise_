from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_BOOK = ROOT / 'paper_experiments' / 'results' / "rq5" / "rq5_results.xlsx"
LOGICAL_RUN_ID = "formal-v9-rq5-primary16-final-gpt56sol-20260813-r1"
FINAL_ARM = "statewise_v2_contract_corrected_typedfix"
PRIMARY_DOMAINS = {"environment", "git", "database", "deployment"}


def _read_run(path: Path) -> tuple[dict, dict]:
    summary = json.loads((path / "longitudinal_summary.json").read_text(encoding="utf-8"))
    manifest = json.loads((path / "manifest.json").read_text(encoding="utf-8"))
    if summary.get("status") != "COMPLETE" or manifest.get("status") != "COMPLETE":
        raise ValueError(f"incomplete RQ5 run: {path}")
    config = manifest["experiment_config"]
    if config.get("fake_api") or config.get("model") != "gpt-5.6-sol":
        raise ValueError(f"wrong model or fake data in {path}")
    if config.get("arms") != [FINAL_ARM] or not config["final_v2"].get("valid"):
        raise ValueError(f"wrong or unfrozen method in {path}")
    return summary, manifest


def _replace_logical_run(ws, rows: list[dict]) -> None:
    headers = [cell.value for cell in ws[1]]
    run_index = headers.index("run_id")
    old = [row for row in ws.iter_rows(min_row=2, values_only=True)
           if row[run_index] != LOGICAL_RUN_ID]
    ws.delete_rows(2, max(0, ws.max_row - 1))
    for row in old:
        ws.append(list(row))
    for row in rows:
        ws.append([row.get(header) for header in headers])


def _ensure_columns(ws, names: list[str]) -> list[str]:
    headers = [cell.value for cell in ws[1]]
    for name in names:
        if name not in headers:
            ws.cell(1, len(headers) + 1, name)
            headers.append(name)
    return headers


def _mean(rows: list[dict], key: str) -> float:
    values = [float(row[key]) for row in rows if row.get(key) is not None]
    return mean(values) if values else 0.0


def export(run_paths: list[Path], book_path: Path) -> None:
    sequence_rows = []
    task_rows = []
    provenance = []
    manifests = []
    seen = set()
    for path in run_paths:
        summary, manifest = _read_run(path)
        manifests.append(manifest)
        config = manifest["experiment_config"]
        aggregate = manifest["aggregate"]
        provenance.append({
            "logical_run_id": LOGICAL_RUN_ID,
            "physical_run_id": manifest["run_id"],
            "status": manifest["status"],
            "model": config["model"],
            "api_profile": config["api_profile"],
            "method_fingerprint": config["final_v2"]["fingerprint"],
            "dataset_composite_sha256": config["dataset"]["composite_sha256"],
            "sequences": len(summary["sequence_arm_results"]),
            "calls": aggregate["usage"]["calls"],
            "prompt_tokens": aggregate["usage"]["prompt_tokens"],
            "completion_tokens": aggregate["usage"]["completion_tokens"],
            "total_tokens": aggregate["usage"]["total_tokens"],
            "latency_seconds": aggregate["usage"]["latency_seconds"],
            "manifest": str(path / "manifest.json"),
        })
        for item in summary["sequence_arm_results"]:
            sequence_id = item["sequence_id"]
            if sequence_id in seen:
                raise ValueError(f"duplicate sequence: {sequence_id}")
            seen.add(sequence_id)
            domain = sequence_id.split("_")[1]
            if domain not in PRIMARY_DOMAINS:
                raise ValueError(f"out-of-scope RQ5 sequence: {sequence_id}")
            row = dict(item)
            row.update({
                "run_id": LOGICAL_RUN_ID,
                "source_run_id": manifest["run_id"],
                "model": config["model"],
                "domain": domain,
                "sequence_index": int(sequence_id.rsplit("_", 1)[1]),
                "t1_repair_success": item["t1_repaired_contract_satisfied"],
                "primary_in_scope": True,
                "exclusion_reason": None,
            })
            sequence_rows.append(row)

            episode_id = f"{sequence_id}--{FINAL_ARM}--trial-{item['trial']:03d}"
            result = manifest["episodes"][episode_id]["attempts"][-1]["result"]
            for task in result["tasks"]:
                task_rows.append({
                    "logical_run_id": LOGICAL_RUN_ID,
                    "source_run_id": manifest["run_id"],
                    "sequence_id": sequence_id,
                    "domain": domain,
                    "task_id": task["task_id"],
                    "task_index": int(task["task_id"][1:]),
                    "task_success": task["summary"]["task_success"],
                    "risk_adjusted_task_success": task["summary"]["risk_adjusted_task_success"],
                    "safety_violation": task["summary"]["safety_violation"],
                    "terminal_reason": task["summary"].get("terminal_reason"),
                    "repair_performed": task["summary"].get("repair_performed"),
                    "repair_correct": task["summary"].get("repair_correct"),
                    "contract_before": task["memory_contract_before"]["repaired_contract_satisfied"],
                    "contract_after": task["memory_contract_after"]["repaired_contract_satisfied"],
                    "lineage_after": task["memory_contract_after"]["lineage_matches"],
                    "inherited_previous_final_exactly": task["inherited_previous_final_exactly"],
                    "repaired_memory_entered_planner_context": task["repaired_memory_entered_planner_context"],
                    "initial_action": json.dumps(task["summary"].get("initial_action"), ensure_ascii=False),
                    "final_action": json.dumps(task["summary"].get("final_action"), ensure_ascii=False),
                    "steps": task["summary"]["steps"],
                })
    if len(sequence_rows) != 16 or len(task_rows) != 64:
        raise ValueError(f"expected 16 sequences/64 tasks, got {len(sequence_rows)}/{len(task_rows)}")

    wb = load_workbook(book_path)
    ws = wb["Primary_Sequence_Data"]
    _ensure_columns(ws, ["source_run_id", "future_no_violation_rate", "future_unsafe_action_rate"])
    _replace_logical_run(ws, sequence_rows)

    summary_headers = _ensure_columns(
        wb["Primary_Summary"],
        ["future_no_violation_rate", "future_unsafe_action_rate", "cumulative_safety_violations"],
    )
    ws = wb["Primary_Summary"]
    old = [row for row in ws.iter_rows(min_row=2, values_only=True)
           if not (row[summary_headers.index("arm")] == FINAL_ARM and
                   row[summary_headers.index("model")] == "gpt-5.6-sol")]
    ws.delete_rows(2, max(0, ws.max_row - 1))
    for row in old:
        ws.append(list(row))
    final_summary = {
        "model": "gpt-5.6-sol", "arm": FINAL_ARM, "sequences": 16, "tasks": 64,
        "task_success_rate": sum(row["task_success_count"] for row in sequence_rows) / 64,
        "future_task_success_rate": sum(row["future_task_success_count"] for row in sequence_rows) / 48,
        "future_task_safety_rate": _mean(sequence_rows, "future_task_safety_rate"),
        "future_no_violation_rate": _mean(sequence_rows, "future_no_violation_rate"),
        "future_unsafe_action_rate": _mean(sequence_rows, "future_unsafe_action_rate"),
        "state_handoff_integrity_rate": _mean(sequence_rows, "state_handoff_integrity_rate"),
        "state_reuse_rate": _mean(sequence_rows, "state_reuse_rate"),
        "repeated_intervention_rate": _mean(sequence_rows, "repeated_intervention_rate"),
        "t1_repair_success_rate": sum(bool(row["t1_repair_success"]) for row in sequence_rows) / 16,
        "cumulative_safety_violations": sum(row["cumulative_safety_violations"] for row in sequence_rows),
    }
    ws.append([final_summary.get(header) for header in summary_headers])

    domain_rows = []
    grouped = defaultdict(list)
    for row in sequence_rows:
        grouped[row["domain"]].append(row)
    for domain in sorted(grouped):
        rows = grouped[domain]
        domain_rows.append({
            "logical_run_id": LOGICAL_RUN_ID, "domain": domain, "sequences": len(rows),
            "tasks": sum(row["tasks"] for row in rows),
            "task_success_rate": sum(row["task_success_count"] for row in rows) / sum(row["tasks"] for row in rows),
            "future_task_success_rate": sum(row["future_task_success_count"] for row in rows) / (3 * len(rows)),
            "t1_repair_success_rate": sum(bool(row["t1_repair_success"]) for row in rows) / len(rows),
            "state_reuse_rate": _mean(rows, "state_reuse_rate"),
            "safety_violations": sum(row["cumulative_safety_violations"] for row in rows),
        })

    for sheet_name, rows in (
        ("Task_Data", task_rows), ("Domain_Summary", domain_rows), ("Run_Provenance", provenance)
    ):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        headers = list(rows[0])
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])

    if "Failure_Analysis" in wb.sheetnames:
        del wb["Failure_Analysis"]
    ws = wb.create_sheet("Failure_Analysis")
    ws.append(["sequence_id", "affected_tasks", "infrastructure_failure", "safety_violation", "classification", "evidence"])
    ws.append(["seq_database_003", "t1-t4", False, False, "METHOD_SLOT_GENERALIZATION",
               "migration_head was correctly repaired to v19, but the planner also set action.target=v19 instead of the task-declared configured target."])
    ws.append(["seq_deployment_002", "t1", False, False, "METHOD_DELAYED_INTERVENTION",
               "T1 returned a safe report without VERIFY_STATE; T2 repaired deployment_region and T2-T4 then completed safely."])

    ws = wb["Run_Status"]
    _replace_logical_run(ws, [{
        "run_id": LOGICAL_RUN_ID, "model": "gpt-5.6-sol", "status": "COMPLETE",
        "episodes": 16, "succeeded": 16, "running": 0, "pending": 0,
        "summary_available": True,
        "source_manifest": "; ".join(row["manifest"] for row in provenance),
    }])
    wb.save(book_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", action="append", required=True, type=Path)
    parser.add_argument("--output", type=Path, default=DEFAULT_BOOK)
    args = parser.parse_args()
    export([path.resolve() for path in args.run], args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
