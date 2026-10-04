from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from scaffold.runners.export_paper_xlsx import _aggregate, _flatten_episode, _load_rows
from scaffold.runners.finalize_formal_v9_rq1_rq3 import (
    exact_mcnemar,
    holm_adjust,
    paired_bootstrap_ci,
)


EXPECTED_DATASET_ID = "w1_contract_corrected_v10_permission_enabled_sensitivity"
EXPECTED_DATASET_SHA = "eefddd2ec3643136bf6bba0ecb395275e690212f55a07c028a39b8e789daed79"
EXPECTED_METHOD_FP = "sha256:0804b97693cb2f4cab236f325347b3c227c391572aa7e3048f85657640cc899e"
STATEWISE = "statewise_v2_contract_corrected_typedfix"
ORACLE = "correct"
ARMS = (
    "correct",
    "corrupted",
    "no_memory",
    "bare_approval",
    "rationale_approval",
    "generic_clarification",
    "memlineage_trust",
    "safety_sentry",
    "ask_or_assume",
    STATEWISE,
)


def _records(path: Path, sheet: str) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = workbook[sheet].iter_rows(values_only=True)
    headers = [str(value) for value in next(values)]
    return headers, [dict(zip(headers, row)) for row in values]


def _write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet.append(["no_rows"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="35546D")


def finalize(
    run_dir: Path,
    overlay_snapshot: Path,
    formal_xlsx: Path,
    formal_recovery_run_dirs: list[Path],
    output_xlsx: Path,
    output_audit: Path,
    expected_run_id: str,
) -> dict[str, Any]:
    manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
    overlay = json.loads(overlay_snapshot.read_text(encoding="utf-8"))
    cases_root = overlay_snapshot.parent / "cases"
    cases = {
        entry["case_id"]: json.loads(
            (cases_root / entry["file"]).read_text(encoding="utf-8")
        )
        for entry in overlay["cases"]
    }
    cohort = set(overlay["sensitivity_cohort_case_ids"])

    matched_rows: list[dict[str, Any]] = []
    statuses = Counter()
    for episode in manifest["episodes"].values():
        statuses[episode["status"]] += 1
        case_id = episode["spec"]["case_id"]
        matched_rows.append(
            _flatten_episode(manifest["run_id"], manifest, episode, cases[case_id])
        )
    matched_rows.sort(key=lambda row: (str(row["arm"]), str(row["case_id"])))

    original_headers, original_rows = _records(formal_xlsx, "Episode_Data")
    original_rows = [row for row in original_rows if row.get("arm") in ARMS]
    formal_terminal_keys = {
        (str(row["arm"]), str(row["case_id"]))
        for row in original_rows
        if row.get("episode_status") != "SUCCEEDED"
    }
    formal_recovery_manifests = []
    formal_replacements: dict[tuple[str, str], dict[str, Any]] = {}
    for recovery_dir in formal_recovery_run_dirs:
        recovery_manifest, recovery_rows = _load_rows(recovery_dir)
        formal_recovery_manifests.append(recovery_manifest)
        for row in recovery_rows:
            key = (str(row["arm"]), str(row["case_id"]))
            if key in formal_terminal_keys and row.get("episode_status") == "SUCCEEDED":
                formal_replacements[key] = row
    original_rows = [
        {
            header: formal_replacements[(str(row["arm"]), str(row["case_id"]))].get(header)
            for header in original_headers
        }
        if (str(row["arm"]), str(row["case_id"])) in formal_replacements
        else row
        for row in original_rows
    ]
    replacements = {(row["arm"], row["case_id"]): row for row in matched_rows}
    combined_rows = []
    before_after = []
    for original in original_rows:
        key = (original["arm"], original["case_id"])
        replacement = replacements.get(key)
        if replacement is None:
            row = dict(original)
            row["permission_policy_source"] = "formal_v9_unchanged"
        else:
            row = {header: replacement.get(header) for header in original_headers}
            row["permission_policy_source"] = "v10_matched_correction"
            before_after.append({
                "arm": key[0],
                "case_id": key[1],
                "before_correct": bool(original.get("behaviorally_correct")),
                "after_correct": bool(replacement.get("behaviorally_correct")),
                "before_action_completed": bool(original.get("action_completed")),
                "after_action_completed": bool(replacement.get("action_completed")),
                "before_violation": bool(original.get("safety_violation")),
                "after_violation": bool(replacement.get("safety_violation")),
            })
        combined_rows.append(row)

    arm_summary = _aggregate(combined_rows, ("arm",))
    domain_summary = _aggregate(combined_rows, ("arm", "domain"))

    outcomes = {
        (row["arm"], row["case_id"]): bool(row["behaviorally_correct"])
        for row in combined_rows
        if row.get("episode_status") == "SUCCEEDED"
    }
    contrast_specs = [
        ("RQ1", "correct", "corrupted"),
        ("RQ1", "no_memory", "corrupted"),
        *[
            ("RQ3", STATEWISE, comparator)
            for comparator in (
                "correct",
                "corrupted",
                "no_memory",
                "memlineage_trust",
                "safety_sentry",
                "ask_or_assume",
            )
        ],
    ]
    paired_contrasts: list[dict[str, Any]] = []
    for rq, focal, comparator in contrast_specs:
        case_ids = sorted(
            {case_id for arm, case_id in outcomes if arm == focal}
            & {case_id for arm, case_id in outcomes if arm == comparator}
        )
        focal_values = [outcomes[(focal, case_id)] for case_id in case_ids]
        comparator_values = [outcomes[(comparator, case_id)] for case_id in case_ids]
        focal_only = sum(
            focal_value and not comparator_value
            for focal_value, comparator_value in zip(focal_values, comparator_values)
        )
        comparator_only = sum(
            comparator_value and not focal_value
            for focal_value, comparator_value in zip(focal_values, comparator_values)
        )
        ci_low, ci_high = paired_bootstrap_ci(focal_values, comparator_values)
        paired_contrasts.append({
            "rq": rq,
            "focal": focal,
            "comparator": comparator,
            "paired_cases": len(case_ids),
            "focal_rate": sum(focal_values) / len(case_ids),
            "comparator_rate": sum(comparator_values) / len(case_ids),
            "risk_difference": (
                sum(focal_values) - sum(comparator_values)
            ) / len(case_ids),
            "bootstrap_ci95_low": ci_low,
            "bootstrap_ci95_high": ci_high,
            "focal_only": focal_only,
            "comparator_only": comparator_only,
            "exact_mcnemar_p": exact_mcnemar(focal_only, comparator_only),
        })
    for rq in ("RQ1", "RQ3"):
        group = [row for row in paired_contrasts if row["rq"] == rq]
        adjusted = holm_adjust([float(row["exact_mcnemar_p"]) for row in group])
        for row, p_value in zip(group, adjusted):
            row["holm_adjusted_p"] = p_value
    by_arm = {row["arm"]: row for row in arm_summary}
    statewise = by_arm.get(STATEWISE, {})
    oracle = by_arm.get(ORACLE, {})
    deployable = [row for row in arm_summary if row["arm"] != ORACLE]
    deployable_best = max(
        (float(row.get("behavioral_correctness_rate") or 0.0) for row in deployable),
        default=0.0,
    )
    statewise_rate = float(statewise.get("behavioral_correctness_rate") or 0.0)
    oracle_rate = float(oracle.get("behavioral_correctness_rate") or 0.0)
    statewise_violation = float(statewise.get("safety_violation_rate") or 0.0)

    config = manifest.get("experiment_config", {})
    dataset = manifest.get("dataset", {})
    implementation = (
        config.get("arm_registry", {}).get(STATEWISE, {}).get("implementation", {})
    )
    failed_api_calls = sum(int(row.get("failed_api_calls") or 0) for row in matched_rows)
    checks = {
        "run_id_matches": manifest.get("run_id") == expected_run_id,
        "run_complete": manifest.get("status") == "COMPLETE",
        "all_200_episodes_succeeded": statuses == Counter({"SUCCEEDED": 200}),
        "arms_match": tuple(config.get("arms", [])) == ARMS,
        "one_trial": config.get("trials") == 1,
        "selected_cases_match_cohort": set(config.get("selected_cases", [])) == cohort,
        "dataset_id_matches": dataset.get("snapshot_id") == EXPECTED_DATASET_ID,
        "dataset_sha_matches": dataset.get("snapshot_sha256") == EXPECTED_DATASET_SHA,
        "method_fingerprint_matches": implementation.get("fingerprint") == EXPECTED_METHOD_FP,
        "process_isolated": config.get("worker_isolation") == "process",
        "live_gpt56sol": config.get("model") == "gpt-5.6-sol" and config.get("fake_api") is False,
        "temperature_zero": float(config.get("temperature")) == 0.0,
        "no_api_failures": failed_api_calls == 0,
        "replacement_pairing_complete": len(replacements) == len(ARMS) * len(cohort),
        "combined_matrix_complete": len(combined_rows) == len(ARMS) * 150,
        "combined_matrix_all_succeeded": all(
            row.get("episode_status") == "SUCCEEDED" for row in combined_rows
        ),
        "all_formal_terminal_failures_recovered": formal_terminal_keys == set(formal_replacements),
        "formal_recovery_provenance_matches": all(
            recovery.get("dataset", {}).get("snapshot_id") == "w1_contract_corrected_v9"
            and recovery.get("experiment_config", {}).get("model") == "gpt-5.6-sol"
            and float(recovery.get("experiment_config", {}).get("temperature")) == 0.0
            and recovery.get("experiment_config", {}).get("worker_isolation") == "process"
            for recovery in formal_recovery_manifests
        ),
        "statewise_zero_violations": statewise_violation == 0.0,
        "statewise_best_deployable": statewise_rate >= deployable_best,
        "statewise_not_materially_below_oracle": statewise_rate + 0.03 >= oracle_rate,
    }
    valid = all(checks.values())

    paper_summary = []
    for arm in ARMS:
        row = by_arm[arm]
        paper_summary.append({
            "arm": arm,
            "objective_correctness": row["behavioral_correctness_rate"],
            "action_completion": row["action_completion_rate"],
            "safety_violation": row["safety_violation_rate"],
            "mean_api_calls": row["mean_api_calls"],
            "mean_total_tokens": row["mean_total_tokens"],
        })

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Paper_Summary", paper_summary)
    _write_sheet(workbook, "Arm_Summary", arm_summary)
    _write_sheet(workbook, "Domain_Summary", domain_summary)
    _write_sheet(workbook, "Paired_Contrasts", paired_contrasts)
    _write_sheet(workbook, "Matched_Cohort", matched_rows)
    _write_sheet(workbook, "Before_After", before_after)
    _write_sheet(workbook, "Combined_Episodes", combined_rows)
    _write_sheet(workbook, "Validation", [
        {"check": key, "passed": value} for key, value in checks.items()
    ])
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)

    audit = {
        "schema_version": "permission-matched-v10-audit-v1",
        "status": "VALID" if valid else "INVALID",
        "replacement_recommended": valid,
        "checks": checks,
        "episode_statuses": dict(statuses),
        "summary": paper_summary,
        "paired_contrasts": paired_contrasts,
        "statewise_vs_oracle_pp": round((statewise_rate - oracle_rate) * 100, 4),
        "statewise_vs_best_deployable_pp": round((statewise_rate - max(
            float(row["objective_correctness"])
            for row in paper_summary if row["arm"] not in {STATEWISE, ORACLE}
        )) * 100, 4),
        "provenance": {
            "run_id": manifest.get("run_id"),
            "dataset_snapshot_id": dataset.get("snapshot_id"),
            "dataset_snapshot_sha256": dataset.get("snapshot_sha256"),
            "method_fingerprint": implementation.get("fingerprint"),
            "model": config.get("model"),
            "temperature": config.get("temperature"),
            "worker_isolation": config.get("worker_isolation"),
            "formal_recovery_run_ids": [
                item.get("run_id") for item in formal_recovery_manifests
            ],
        },
    }
    output_audit.parent.mkdir(parents=True, exist_ok=True)
    output_audit.write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return audit


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--overlay-snapshot", type=Path, required=True)
    parser.add_argument("--formal-xlsx", type=Path, required=True)
    parser.add_argument("--formal-recovery-run-dir", type=Path, action="append", default=[])
    parser.add_argument("--output-xlsx", type=Path, required=True)
    parser.add_argument("--output-audit", type=Path, required=True)
    parser.add_argument("--expected-run-id", required=True)
    args = parser.parse_args()
    audit = finalize(
        args.run_dir.resolve(),
        args.overlay_snapshot.resolve(),
        args.formal_xlsx.resolve(),
        [path.resolve() for path in args.formal_recovery_run_dir],
        args.output_xlsx.resolve(),
        args.output_audit.resolve(),
        args.expected_run_id,
    )
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0 if audit["status"] == "VALID" else 1


if __name__ == "__main__":
    raise SystemExit(main())
