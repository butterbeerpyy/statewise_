import json
import pathlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = pathlib.Path(__file__).resolve().parents[2]
SOURCE = ROOT / "runs" / "rq5_primary_scope.json"
OUT = ROOT / "runs" / "rq5_primary_scope.xlsx"

data = json.loads(SOURCE.read_text(encoding="utf-8"))
cells = data["cells"]

wb = openpyxl.Workbook()
wb.remove(wb.active)


ws1 = wb.create_sheet("primary_scope")
ws1.append([
    "arm", "model", "n_sequences", "task_success_rate", "future_task_success_rate",
    "future_task_safety_rate", "state_handoff_integrity_rate", "state_reuse_rate",
    "t1_repaired_contract_satisfied", "cumulative_safety_violations"
])

for cell in sorted(cells, key=lambda c: (c["arm"] != "statewise_v2_retrieval", c["arm"], c["model"])):
    p = cell["primary_scope"]
    if p.get("sequences", 0) > 0:
        ws1.append([
            cell["arm"], cell["model"], p.get("sequences"),
            p.get("task_success_rate"), p.get("future_task_success_rate"),
            p.get("future_task_safety_rate"), p.get("state_handoff_integrity_rate"),
            p.get("state_reuse_rate"), p.get("t1_repaired_contract_satisfied"),
            p.get("cumulative_safety_violations"),
        ])


ws2 = wb.create_sheet("filesystem_only")
ws2.append([
    "arm", "model", "n_sequences", "task_success_rate", "future_task_success_rate",
    "future_task_safety_rate", "state_handoff_integrity_rate", "state_reuse_rate",
    "t1_repaired_contract_satisfied", "cumulative_safety_violations"
])

for cell in sorted(cells, key=lambda c: (c["arm"] != "statewise_v2_retrieval", c["arm"], c["model"])):
    f = cell["filesystem_only"]
    if f.get("sequences", 0) > 0:
        ws2.append([
            cell["arm"], cell["model"], f.get("sequences"),
            f.get("task_success_rate"), f.get("future_task_success_rate"),
            f.get("future_task_safety_rate"), f.get("state_handoff_integrity_rate"),
            f.get("state_reuse_rate"), f.get("t1_repaired_contract_satisfied"),
            f.get("cumulative_safety_violations"),
        ])


header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
for ws in (ws1, ws2):
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

wb.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)}")
